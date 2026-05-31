"""
Generates practical_work_18_full.docx
Practical Work 18 – expanded step-by-step solution with all formulas.
"""

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from scipy import stats
from scipy.integrate import trapezoid as scipy_trapz
import warnings, os
warnings.filterwarnings('ignore')

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ─────────────────────────────────────────────────
#  PRIMARY DATA
# ─────────────────────────────────────────────────
names = [
    'Андрієнко','Бойко','Василенко','Гнатенко','Демченко',
    'Єременко','Захаренко','Іваненко','Коваленко','Литвиненко',
    'Марченко','Назаренко','Олексієнко','Петренко','Романенко',
    'Савченко','Тимченко','Удовенко','Філіпенко','Харченко'
]
body_weight = np.array([75,82,68,90,78,85,72,88,76,81,
                         70,93,65,87,79,84,73,91,77,83])
grip_1 = np.array([48,55,42,61,51,58,46,63,49,54,
                    44,67,40,59,52,57,47,65,50,56])
grip_2 = np.array([47,56,43,60,52,59,45,62,50,55,
                    45,66,41,60,53,58,48,64,51,55])
grip_best = np.maximum(grip_1, grip_2)
rel_str   = grip_best / body_weight
n = len(grip_best)

# ── DYNAMOGRAM: 15 uniformly-spaced points, Δt = 25 ms ──
t_ms  = np.array([  0,  25,  50,  75, 100, 125, 150, 175,
                   200, 225, 250, 275, 300, 325, 350], dtype=float)
F_dyn = np.array([785,1010,1400,1830,2130,2200,1880,1570,
                  1280,1090, 975, 905, 855, 820, 795], dtype=float)
t_s   = t_ms / 1000.0          # seconds
dt_s  = 0.025                   # Δt = 25 ms = 0.025 s
T_tot = 0.350                   # total time, s
m_body_dyn = 80.0               # kg
F_body     = m_body_dyn * 9.81  # N  ≈ 784.8 N

# Trapezoid rule – show full table
I_trap  = scipy_trapz(F_dyn, t_s)     # total impulse
F_avg_I = I_trap / T_tot              # mean force from impulse
F_max_d = np.max(F_dyn)               # peak force
t_peak  = t_ms[np.argmax(F_dyn)]      # time to peak (ms)
G_F     = F_max_d / (t_peak / 1000.0) # gradient (N/s)

# ─────────────────────────────────────────────────
#  STATISTICS  (best attempt)
# ─────────────────────────────────────────────────
M1  = np.mean(grip_1);   s1  = np.std(grip_1,  ddof=1)
M2  = np.mean(grip_2);   s2  = np.std(grip_2,  ddof=1)
Mb  = np.mean(grip_best);sb  = np.std(grip_best,ddof=1)
Mr  = np.mean(rel_str);  sr  = np.std(rel_str,  ddof=1)

V1  = s1/M1*100;   V2=s2/M2*100;   Vb=sb/Mb*100;   Vr=sr/Mr*100
m1  = s1/np.sqrt(n); m2=s2/np.sqrt(n); mb=sb/np.sqrt(n)

# Pearson r
dev1 = grip_1 - M1;   dev2 = grip_2 - M2
num_r  = np.sum(dev1*dev2)
den_r  = np.sqrt(np.sum(dev1**2)*np.sum(dev2**2))
r_tt   = num_r / den_r
t_obs  = r_tt*np.sqrt(n-2)/np.sqrt(1-r_tt**2)
t_cr   = stats.t.ppf(0.975, df=n-2)
p_val  = stats.pearsonr(grip_1,grip_2)[1]
SEM    = sb * np.sqrt(1-r_tt)
rel_level = 'ВІДМІННА' if r_tt >= 0.90 else ('ХОРОША' if r_tt >= 0.80 else 'ЗАДОВІЛЬНА')

def score5(x, m, s):
    if   x < m-1.5*s: return 1
    elif x < m-0.5*s: return 2
    elif x < m+0.5*s: return 3
    elif x < m+1.5*s: return 4
    else:              return 5

level_lbl={1:'Низький',2:'Нижче середнього',3:'Середній',
           4:'Вище середнього',5:'Високий'}
scores_b = [score5(v,Mb,sb) for v in grip_best]
scores_r = [score5(v,Mr,sr) for v in rel_str]

# ─────────────────────────────────────────────────
#  GENERATE GRAPHS
# ─────────────────────────────────────────────────
plt.rcParams.update({'font.family':'DejaVu Sans','axes.grid':True,'grid.alpha':0.3})

def save_g1():
    fig,ax=plt.subplots(figsize=(13,6))
    ax.plot(t_ms,F_dyn,'o-',color='royalblue',lw=2,ms=6,
            label='F(t) – сила реакції опори (Н)')
    # shade area = impulse
    ax.fill_between(t_ms,F_dyn,alpha=0.13,color='royalblue',
                    label=f'Імпульс I = {I_trap:.1f} Н·с')
    ax.axhline(F_body,color='gray',ls='--',lw=1.3,
               label=f'Вага тіла = {F_body:.0f} Н')
    ax.axhline(F_avg_I,color='orange',ls='-.',lw=1.8,
               label=f'Середня сила F̄ = {F_avg_I:.0f} Н')
    # peak annotation
    ax.annotate(f'F_max = {F_max_d:.0f} Н\nt = {t_peak:.0f} мс',
                xy=(t_peak,F_max_d),xytext=(t_peak+35,F_max_d-300),
                arrowprops=dict(arrowstyle='->',color='red',lw=1.8),
                fontsize=11,color='red',
                bbox=dict(boxstyle='round,pad=0.3',fc='lightyellow',ec='red'))
    # gradient arrow
    ax.annotate('',xy=(t_peak,F_max_d),xytext=(0,F_body),
                arrowprops=dict(arrowstyle='->',color='green',lw=1.8,linestyle='dashed'))
    ax.text(t_peak/2-10,(F_max_d+F_body)/2+30,
            f'G_F = {G_F:.0f} Н/с',color='green',fontsize=10)
    ax.set_xlabel('Час, мс',fontsize=13); ax.set_ylabel('Сила, Н',fontsize=13)
    ax.set_title('Графік 1. Динамограма відштовхування спортсмена від опори',
                 fontsize=13,fontweight='bold')
    ax.legend(loc='upper right',fontsize=10)
    ax.set_xlim(-5,360); ax.set_ylim(0,F_max_d*1.18)
    plt.tight_layout(); plt.savefig('g1.png',dpi=150,bbox_inches='tight'); plt.close()

def save_g2():
    fig,(ax1,ax2)=plt.subplots(1,2,figsize=(14,6))
    bins=np.arange(38,72,4)
    counts,_,_=ax1.hist(grip_best,bins=bins,edgecolor='white',
                         lw=1.5,color='steelblue',alpha=0.75)
    xr=np.linspace(28,76,300)
    ax1.plot(xr,stats.norm.pdf(xr,Mb,sb)*n*(bins[1]-bins[0]),
             'navy',lw=2.5,label='Нормальний розподіл')
    sc_cols=['#ef5350','#ff9800','#66bb6a','#42a5f5','#ab47bc']
    bounds_v=[Mb-1.5*sb,Mb-0.5*sb,Mb,Mb+0.5*sb,Mb+1.5*sb]
    labels_v=['M−1.5σ','M−0.5σ','M','M+0.5σ','M+1.5σ']
    for b,col,lbl in zip(bounds_v,sc_cols,labels_v):
        ax1.axvline(b,color=col,ls='--',lw=1.8)
        ax1.text(b,counts.max()*1.07,lbl,ha='center',fontsize=9,color=col)
    ax1.set_xlabel('Сила кисті, кгс',fontsize=12)
    ax1.set_ylabel('Кількість спортсменів',fontsize=12)
    ax1.set_title('Розподіл результатів та межі нормувальної шкали',fontsize=11,fontweight='bold')
    ax1.legend(fontsize=10)
    lo_b=[0,Mb-1.5*sb,Mb-0.5*sb,Mb+0.5*sb,Mb+1.5*sb]
    hi_b=[Mb-1.5*sb,Mb-0.5*sb,Mb+0.5*sb,Mb+1.5*sb,80]
    nl=['Низький','Нижче\nсередн.','Середній','Вище\nсередн.','Високий']
    for i,(lo,hi,col,nm) in enumerate(zip(lo_b,hi_b,sc_cols,nl)):
        ax2.barh(i,hi-lo,left=lo,color=col,alpha=0.75,edgecolor='white',height=0.7)
        ax2.text((lo+hi)/2,i,nm,ha='center',va='center',
                 fontsize=11,fontweight='bold',color='white')
        ax2.text(lo+0.3,i-0.42,f'{lo:.1f}',ha='left',fontsize=9)
        if hi<80: ax2.text(hi-0.3,i-0.42,f'{hi:.1f}',ha='right',fontsize=9)
    for val in grip_best:
        lvl=score5(val,Mb,sb)-1
        ax2.plot(val,lvl,'k|',ms=16,mew=2,alpha=0.6)
    ax2.set_xlabel('Сила кисті, кгс',fontsize=12)
    ax2.set_yticks(range(5))
    ax2.set_yticklabels(['1–Низький','2–Нижче серед.','3–Середній',
                          '4–Вище серед.','5–Високий'],fontsize=10)
    ax2.set_title('Нормувальна шкала',fontsize=11,fontweight='bold')
    ax2.set_xlim(30,75)
    plt.tight_layout(); plt.savefig('g2.png',dpi=150,bbox_inches='tight'); plt.close()

def save_g3():
    fig,ax=plt.subplots(figsize=(8,8))
    ax.scatter(grip_1,grip_2,s=90,color='royalblue',edgecolors='navy',alpha=0.8,zorder=5)
    slope,intercept,*_=stats.linregress(grip_1,grip_2)
    xl=np.linspace(37,70,100)
    ax.plot(xl,slope*xl+intercept,'r-',lw=2,
            label=f'y = {slope:.3f}x + {intercept:.3f}')
    lim=min(grip_1.min(),grip_2.min())-2
    liM=max(grip_1.max(),grip_2.max())+2
    ax.plot([lim,liM],[lim,liM],'g--',lw=1.5,alpha=0.5,label='Бісектриса')
    for i,nm in enumerate(names):
        ax.annotate(nm[:4],(grip_1[i],grip_2[i]),
                    xytext=(3,3),textcoords='offset points',fontsize=7,color='gray')
    ax.set_xlabel('Спроба 1, кгс',fontsize=13); ax.set_ylabel('Спроба 2, кгс',fontsize=13)
    ax.set_title(f'Графік 3. Кореляційне поле (r_tt = {r_tt:.4f})',
                 fontsize=12,fontweight='bold')
    ax.legend(fontsize=10); ax.set_aspect('equal')
    ax.set_xlim(lim,liM); ax.set_ylim(lim,liM)
    plt.tight_layout(); plt.savefig('g3.png',dpi=150,bbox_inches='tight'); plt.close()

def save_g4():
    fig,axes=plt.subplots(1,3,figsize=(18,6))
    cols=[plt.cm.RdYlGn(score5(v,Mb,sb)/5) for v in grip_best]
    axes[0].bar(range(1,n+1),grip_best,color=cols,edgecolor='white',lw=0.8)
    axes[0].axhline(Mb,color='navy',ls='--',lw=2,label=f'M={Mb:.1f}')
    axes[0].axhline(Mb+sb,color='green',ls=':',lw=1.5,label=f'M+σ={Mb+sb:.1f}')
    axes[0].axhline(Mb-sb,color='orange',ls=':',lw=1.5,label=f'M-σ={Mb-sb:.1f}')
    axes[0].set_xticks(range(1,n+1))
    axes[0].set_xticklabels([nm[:4] for nm in names],rotation=45,ha='right',fontsize=8)
    axes[0].set_ylabel('Сила кисті, кгс',fontsize=11)
    axes[0].set_title('Абсолютна сила кисті',fontsize=12,fontweight='bold')
    axes[0].legend(fontsize=9)
    cr=[plt.cm.RdYlGn(score5(v,Mr,sr)/5) for v in rel_str]
    axes[1].bar(range(1,n+1),rel_str,color=cr,edgecolor='white',lw=0.8)
    axes[1].axhline(Mr,color='navy',ls='--',lw=2,label=f'M={Mr:.3f}')
    axes[1].set_xticks(range(1,n+1))
    axes[1].set_xticklabels([nm[:4] for nm in names],rotation=45,ha='right',fontsize=8)
    axes[1].set_ylabel('кгс/кг',fontsize=11)
    axes[1].set_title('Відносна сила',fontsize=12,fontweight='bold')
    axes[1].legend(fontsize=9)
    from collections import Counter
    cnt=Counter([level_lbl[s] for s in scores_b])
    order=['Низький','Нижче середнього','Середній','Вище середнього','Високий']
    sc_c=['#ef5350','#ff9800','#66bb6a','#42a5f5','#ab47bc']
    sz=[cnt.get(l,0) for l in order]; nz=[(s,l,c) for s,l,c in zip(sz,order,sc_c) if s>0]
    sv,lv,cv=zip(*nz)
    axes[2].pie(sv,labels=lv,colors=cv,
                autopct=lambda p:f'{p:.0f}%\n({int(round(p*n/100))} ос.)',
                startangle=90,pctdistance=0.75,
                wedgeprops=dict(edgecolor='white',linewidth=2))
    axes[2].set_title('Розподіл за рівнями\nпідготовленості',fontsize=12,fontweight='bold')
    plt.suptitle('Графік 4. Порівняльний аналіз силової підготовленості',
                 fontsize=13,fontweight='bold',y=1.02)
    plt.tight_layout(); plt.savefig('g4.png',dpi=150,bbox_inches='tight'); plt.close()

print("Generating graphs…")
save_g1(); save_g2(); save_g3(); save_g4()
print("Graphs saved.")

# ─────────────────────────────────────────────────
#  DOCX HELPERS
# ─────────────────────────────────────────────────
def new_doc():
    doc=Document()
    s=doc.styles['Normal']
    s.font.name='Times New Roman'; s.font.size=Pt(14)
    s.paragraph_format.space_after=Pt(0)
    for sec in doc.sections:
        sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
        sec.left_margin=Cm(3); sec.right_margin=Cm(1.5)
    return doc

def h1(doc,text):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(8)
    r=p.add_run(text); r.font.name='Times New Roman'
    r.font.size=Pt(15); r.font.bold=True
    return p

def h2(doc,text):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.font.name='Times New Roman'
    r.font.size=Pt(14); r.font.bold=True
    return p

def para(doc,text,indent=True,bold=False,italic=False,align=WD_ALIGN_PARAGRAPH.JUSTIFY,size=14):
    p=doc.add_paragraph()
    p.alignment=align
    p.paragraph_format.first_line_indent=Cm(1.25) if indent else Cm(0)
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.ONE_POINT_FIVE
    r=p.add_run(text); r.font.name='Times New Roman'
    r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    return p

def formula(doc,text):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(text); r.font.name='Times New Roman'
    r.font.size=Pt(13); r.font.italic=True
    return p

def caption(doc,text):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(text); r.font.name='Times New Roman'
    r.font.size=Pt(13); r.font.italic=True
    return p

def spacer(doc): doc.add_paragraph()

def set_borders(tbl):
    t=tbl._tbl; pr=t.tblPr if t.tblPr is not None else OxmlElement('w:tblPr')
    bd=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        tag=OxmlElement(f'w:{e}')
        tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),'4')
        tag.set(qn('w:space'),'0'); tag.set(qn('w:color'),'000000')
        bd.append(tag)
    pr.append(bd)

def shade(cell,hex_c):
    pr=cell._tc.get_or_add_tcPr()
    sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto')
    sh.set(qn('w:fill'),hex_c); pr.append(sh)

def cell_text(cell,text,bold=False,center=False,italic=False,size=11,white=False):
    cell.text=''
    p=cell.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(str(text)); r.font.name='Times New Roman'
    r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    if white: r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)

HDR='1F4E79'       # dark blue header
ALT='EBF5FB'       # light blue alt row
GLD='FFD700'       # gold
SLV='DCDCDC'       # silver
BRZ='F4A460'       # bronze
GRN='D5F5E3'       # green highlight
RED='FADBD8'       # red highlight

def header_row(tbl_row,texts,widths,shade_hex=HDR):
    for cell,text,w in zip(tbl_row.cells,texts,widths):
        cell.width=w
        cell_text(cell,text,bold=True,center=True,white=True,size=11)
        shade(cell,shade_hex)

def data_row(tbl_row,values,widths,centers,bg=None,size=11):
    for cell,val,w,c in zip(tbl_row.cells,values,widths,centers):
        cell.width=w
        cell_text(cell,val,center=c,size=size)
        if bg: shade(cell,bg)

def add_image(doc,path,w_cm=15.5):
    if os.path.exists(path):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path,width=Cm(w_cm))

# ─────────────────────────────────────────────────
#  BUILD DOCUMENT
# ─────────────────────────────────────────────────
doc=new_doc()

# ══ TITLE ══════════════════════════════════════════
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("ПРАКТИЧНА РОБОТА № 18")
r.font.name='Times New Roman'; r.font.size=Pt(18); r.font.bold=True

p2=doc.add_paragraph()
p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
r2=p2.add_run("Тема: Метрологічні основи контролю за рівнем\nрозвитку фізичної якості сили")
r2.font.name='Times New Roman'; r2.font.size=Pt(16); r2.font.bold=True
p2.paragraph_format.space_after=Pt(6)

p3=doc.add_paragraph()
p3.alignment=WD_ALIGN_PARAGRAPH.CENTER
r3=p3.add_run("Мета: навчитися методиці оцінки силових якостей спортсменів.")
r3.font.name='Times New Roman'; r3.font.size=Pt(14); r3.font.italic=True
p3.paragraph_format.space_after=Pt(20)

# ══ SECTION 1 ══════════════════════════════════════
h1(doc,"1. ВИЗНАЧЕННЯ СИЛИ ЯК ФІЗИЧНОЇ ЯКОСТІ ЛЮДИНИ")

para(doc,
     "Здатність долати зовнішній опір або протидіяти йому за допомогою м'язових зусиль "
     "називають силовими якостями. Від рівня їх розвитку залежать досягнення у всіх "
     "видах спорту. Перші механічні прилади для вимірювання сили з'явились у XVIII ст.")

h2(doc,"1.1 Показники силових якостей та їх формули")

para(doc,"При контролі за силовими якостями виділяють три групи показників:",indent=True)

rows_sh=[
    ["Група","Показник","Формула","Одиниця"],
    ["Основні","Максимальна сила","F_max = max F(t)","Н (Ньютон)"],
    ["Основні","Середня сила","F̄ = I / T","Н"],
    ["Інтегральні","Імпульс сили","I = ∫ F(t) dt  або  I = F · t (при F = const)","Н · с"],
    ["Диференціальні","Градієнт сили","G_F = dF/dt  або  G_F = F_max / t(F_max)","Н/с"],
]
tsh=doc.add_table(rows=5,cols=4)
tsh.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(tsh)
wsh=[Cm(3.0),Cm(3.5),Cm(7.0),Cm(2.0)]
for i,row in enumerate(rows_sh):
    if i==0:
        header_row(tsh.rows[0],row,wsh)
    else:
        for j,(v,w) in enumerate(zip(row,wsh)):
            tsh.rows[i].cells[j].width=w
            cell_text(tsh.rows[i].cells[j],v,center=(j in (0,3)),size=11)
            if i%2==0: shade(tsh.rows[i].cells[j],ALT)

para(doc,
     "Пояснення до формул:",indent=False,bold=True)
for txt in [
    "Імпульс сили I = ∫F(t) dt — визначається як площа під кривою F(t). "
    "Чисельно розраховується методом трапецій: I ≈ Σ (F[i] + F[i+1])/2 · Δtᵢ",
    "Середня сила F̄ = I/T — умовна постійна сила, яка дає той самий імпульс за час T.",
    "Градієнт сили G_F = F_max / t(F_max) — характеризує вибуховість зусилля (швидкість наростання сили).",
]:
    para(doc, "• " + txt, indent=True)

spacer(doc)

# ══ SECTION 2 ══════════════════════════════════════
h1(doc,"2. ВИДИ СИЛИ")

para(doc,
     "Силові якості не однорідні. Залежно від умов прояву і завдань контролю "
     "виділяють такі основні види:",indent=True)

kinds=[
    ["Вид сили","Характеристика","Вид спорту / вправи"],
    ["Максимальна сила","Найвищий рівень напруги м'яза незалежно від часу","Важка атлетика, боротьба"],
    ["Вибухова сила","Досягнення F_max за найкоротший час (великий G_F)","Стрибки, метання, спринт"],
    ["Силова витривалість","Підтримання значних зусиль протягом тривалого часу","Веслування, плавання, велоспорт"],
    ["Відносна сила","F_abs / маса тіла (кгс/кг)","Єдиноборства, гімнастика"],
    ["Динамічна сила","Сила, що проявляється під час руху","Практично всі циклічні і ациклічні рухи"],
    ["Статична (ізометрична) сила","Сила без руху (без зміни довжини м'яза)","Боротьба, важка атлетика (фаза утримання)"],
]
tk=doc.add_table(rows=len(kinds),cols=3)
tk.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(tk)
wk=[Cm(4.0),Cm(7.5),Cm(5.0)]
for i,row in enumerate(kinds):
    if i==0: header_row(tk.rows[0],row,wk)
    else:
        for j,(v,w) in enumerate(zip(row,wk)):
            tk.rows[i].cells[j].width=w
            cell_text(tk.rows[i].cells[j],v,center=(j==0),size=11)
            if i%2==0: shade(tk.rows[i].cells[j],ALT)

spacer(doc)

# ══ SECTION 3 ══════════════════════════════════════
h1(doc,"3. КОНТРОЛЬ ЗА РІВНЕМ РОЗВИТКУ СИЛИ")

h2(doc,"3.1 Способи реєстрації та метрологічні вимоги")

for t_ in [
    "1) Без вимірювальної апаратури: оцінка за найбільшою вагою, яку може підняти "
    "спортсмен (прямий спосіб) або за результатом стрибка, метання (непрямий спосіб).",
    "2) З вимірювальними пристроями: а) механічні динамометри пружинного типу (ДПУ: 1000, "
    "2000, 5000 Н; похибка ≤ 2%); б) тензометричні пристрої (найточніші); "
    "в) інерційні динамографи (вимірювання сили в русі).",
]:
    para(doc,"• "+t_,indent=True)

para(doc,"Специфічні метрологічні вимоги при вимірюванні сили:",bold=True,indent=False)
for t_ in [
    "стандартизувати положення тіла та кут у суглобі в повторних спробах;",
    "враховувати довжину сегментів тіла при вимірюванні моментів сили;",
    "враховувати напрям вектора сили відносно осі вимірювального пристрою.",
]:
    para(doc,"• "+t_,indent=True)

h2(doc,"3.2 Надійність силових тестів")

para(doc,
     "Надійність тесту — відтворюваність результату при повторному вимірюванні "
     "в однакових умовах. Оцінюється коефіцієнтом ретест-кореляції r_tt.",indent=True)

rel_t=[
    ["Тип пристрою","r_tt","Рівень надійності"],
    ["Механічні динамометри (пружинні)","0,60 – 0,80","Задовільний – хороший"],
    ["Градієнт сили (будь-який пристрій)","0,70 – 0,80","Задовільний"],
    ["Тензометричні пристрої (F_max)","0,85 – 0,95","Хороший – відмінний"],
]
trl=doc.add_table(rows=len(rel_t),cols=3)
trl.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(trl)
wrl=[Cm(7.0),Cm(3.0),Cm(4.5)]
for i,row in enumerate(rel_t):
    if i==0: header_row(trl.rows[0],row,wrl)
    else:
        for j,(v,w) in enumerate(zip(row,wrl)):
            trl.rows[i].cells[j].width=w
            cell_text(trl.rows[i].cells[j],v,center=(j==1),size=11)
            if i%2==0: shade(trl.rows[i].cells[j],ALT)

spacer(doc)

# ══ RGR No.2 ═══════════════════════════════════════
h1(doc,"РОЗРАХУНКОВО-ГРАФІЧНА РОБОТА № 2")

p_sub=doc.add_paragraph()
p_sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
rs=p_sub.add_run("«Метрологічний аналіз показників силової підготовленості спортсменів»")
rs.font.name='Times New Roman'; rs.font.size=Pt(14); rs.font.bold=True

spacer(doc)

# ════════════════════════════════════════════════════
#  TASK 1 — PRIMARY DATA
# ════════════════════════════════════════════════════
h2(doc,"Завдання 1. Вихідні дані")

para(doc,
     "Вимірювання максимальної сили стискання правої кисті (динамометрія кисті) "
     "проведено для n = 20 спортсменів за допомогою механічного кистьового "
     "динамометра. Кожен спортсмен виконував по 2 спроби; у подальших розрахунках "
     "використовується краща (максимальна) з двох спроб.",indent=True)

caption(doc,"Таблиця 1. Вихідні дані вимірювань динамометрії кисті")

h_t1=["№","Спортсмен","Маса тіла,\nкг","Спроба 1,\nкгс","Спроба 2,\nкгс","Краща\nспроба, кгс","Відносна\nсила, кгс/кг"]
w_t1=[Cm(0.8),Cm(3.2),Cm(1.9),Cm(1.9),Cm(1.9),Cm(2.1),Cm(2.3)]
t1=doc.add_table(rows=n+2,cols=7)
t1.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(t1)
header_row(t1.rows[0],h_t1,w_t1)
for i in range(n):
    row=t1.rows[i+1]
    bg=ALT if i%2==0 else None
    vals=[i+1,names[i],body_weight[i],grip_1[i],grip_2[i],
          grip_best[i],f"{rel_str[i]:.3f}"]
    data_row(row,vals,w_t1,[True,False,True,True,True,True,True],bg=bg)
# Sum row
sr_row=t1.rows[n+1]
for j,w in enumerate(w_t1): sr_row.cells[j].width=w
shade(sr_row.cells[0],'D6EAF0')
cell_text(sr_row.cells[0],"Σ",bold=True,center=True)
shade(sr_row.cells[1],'D6EAF0')
cell_text(sr_row.cells[1],"Сума / Середнє",bold=True,center=True)
for j,(s_val,lab) in enumerate([(np.sum(body_weight),''),(np.sum(grip_1),''),(np.sum(grip_2),''),
                                  (np.sum(grip_best),''),(f"{np.mean(rel_str):.3f}",'')],start=2):
    sr_row.cells[j].width=w_t1[j]
    cell_text(sr_row.cells[j],s_val,bold=True,center=True)
    shade(sr_row.cells[j],'D6EAF0')

para(doc,
     f"Сума результатів кращої спроби: Σxᵢ = {np.sum(grip_best)} кгс  "
     f"(потрібна для розрахунку середнього значення).",indent=True)

spacer(doc)

# ════════════════════════════════════════════════════
#  TASK 2 — DESCRIPTIVE STATISTICS  (STEP BY STEP)
# ════════════════════════════════════════════════════
h2(doc,"Завдання 2. Розрахунок статистичних характеристик (покроково)")

para(doc,
     "Статистичний аналіз проводимо для кращої спроби. Розраховуємо: середнє "
     "арифметичне (M), середнє квадратичне відхилення (σ), коефіцієнт варіації (V%), "
     "похибку середнього (m).",indent=True)

# ── 2.1 MEAN ──
h2(doc,"2.1 Середнє арифметичне (M)")

formula(doc,"M = (x₁ + x₂ + … + xₙ) / n")

para(doc,
     f"Підставляємо дані (n = {n}). Виписуємо всі значення кращої спроби:",indent=True)

vals_str = " + ".join(str(v) for v in grip_best)
formula(doc, f"M = ({vals_str}) / {n}")
formula(doc, f"M = {np.sum(grip_best)} / {n}")
formula(doc, f"M = {Mb:.2f} кгс")

para(doc,
     f"Отже, середня максимальна сила кисті у групі становить "
     f"M = {Mb:.2f} кгс.",indent=True)

spacer(doc)

# ── 2.2 STANDARD DEVIATION ──
h2(doc,"2.2 Середнє квадратичне відхилення (σ)")

formula(doc,"σ = √[ Σ(xᵢ − M)² / (n − 1) ]")

para(doc,
     "Для розрахунку σ заповнюємо допоміжну таблицю: "
     "знаходимо відхилення кожного значення від середнього (xᵢ − M) "
     "та їх квадрати (xᵢ − M)².",indent=True)

caption(doc,"Таблиця 2. Розрахунок середнього квадратичного відхилення")

dev_b = grip_best - Mb
sq_b  = dev_b**2

h_t2=["№","Спортсмен","xᵢ, кгс","xᵢ − M","(xᵢ − M)²"]
w_t2=[Cm(0.8),Cm(3.5),Cm(2.2),Cm(2.8),Cm(3.2)]
t2=doc.add_table(rows=n+2,cols=5)
t2.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(t2)
header_row(t2.rows[0],h_t2,w_t2)
for i in range(n):
    row=t2.rows[i+1]
    bg=ALT if i%2==0 else None
    vals=[i+1,names[i],grip_best[i],f"{dev_b[i]:.2f}",f"{sq_b[i]:.2f}"]
    data_row(row,vals,w_t2,[True,False,True,True,True],bg=bg)
# totals row
tr2=t2.rows[n+1]
for j,w in enumerate(w_t2): tr2.cells[j].width=w
for j,v in enumerate(["","","","Σ(xᵢ-M)²  =",f"{np.sum(sq_b):.2f}"]):
    cell_text(tr2.cells[j],v,bold=True,center=True)
    shade(tr2.cells[j],'D6EAF0')

formula(doc,
        f"σ = √( Σ(xᵢ−M)² / (n−1) ) = √( {np.sum(sq_b):.2f} / ({n}−1) )")
formula(doc,
        f"σ = √( {np.sum(sq_b):.2f} / {n-1} ) = √{np.sum(sq_b)/(n-1):.4f}")
formula(doc, f"σ = {sb:.2f} кгс")

spacer(doc)

# ── 2.3 COEFFICIENT OF VARIATION ──
h2(doc,"2.3 Коефіцієнт варіації (V%)")

formula(doc,"V = (σ / M) · 100%")
formula(doc,f"V = ({sb:.2f} / {Mb:.2f}) · 100%")
formula(doc,f"V = {sb/Mb:.4f} · 100% = {Vb:.1f}%")

para(doc,
     f"V = {Vb:.1f}% — "
     f"{'однорідна група (V < 10%): результати близькі між собою.' if Vb<10 else 'помірно однорідна група (10% ≤ V < 20%): помірний розкид.' if Vb<20 else 'неоднорідна група (V ≥ 20%): значний розкид.'} "
     f"Висновок: тест можна вважати метрологічно коректним для даної групи.",indent=True)

spacer(doc)

# ── 2.4 STANDARD ERROR ──
h2(doc,"2.4 Похибка середнього (m)")

formula(doc,"m = σ / √n")
formula(doc,f"m = {sb:.2f} / √{n} = {sb:.2f} / {np.sqrt(n):.4f}")
formula(doc,f"m = {mb:.2f} кгс")

para(doc,
     f"Це означає: з імовірністю ≈ 68% справжнє середнє генеральної сукупності "
     f"лежить у діапазоні M ± m = {Mb:.2f} ± {mb:.2f} кгс.",indent=True)

spacer(doc)

# ── SUMMARY STATS TABLE ──
caption(doc,"Таблиця 3. Зведена описова статистика")

h_t3=["Показник","Спроба 1\n(кгс)","Спроба 2\n(кгс)","Краща спроба\n(кгс)","Відносна\nсила (кгс/кг)"]
w_t3=[Cm(5.2),Cm(2.5),Cm(2.5),Cm(2.8),Cm(2.8)]
rows_t3=[
    h_t3,
    ["M (середнє)",f"{M1:.2f}",f"{M2:.2f}",f"{Mb:.2f}",f"{Mr:.3f}"],
    ["σ (відхилення)",f"{s1:.2f}",f"{s2:.2f}",f"{sb:.2f}",f"{sr:.4f}"],
    ["V% (варіація)",f"{V1:.1f}%",f"{V2:.1f}%",f"{Vb:.1f}%",f"{Vr:.1f}%"],
    ["m (похибка середнього)",f"{m1:.2f}",f"{m2:.2f}",f"{mb:.2f}",f"{sr/np.sqrt(n):.4f}"],
    ["Мінімум",str(np.min(grip_1)),str(np.min(grip_2)),str(np.min(grip_best)),f"{np.min(rel_str):.3f}"],
    ["Максимум",str(np.max(grip_1)),str(np.max(grip_2)),str(np.max(grip_best)),f"{np.max(rel_str):.3f}"],
    ["Медіана",f"{np.median(grip_1):.1f}",f"{np.median(grip_2):.1f}",f"{np.median(grip_best):.1f}",f"{np.median(rel_str):.3f}"],
]
t3=doc.add_table(rows=len(rows_t3),cols=5)
t3.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(t3)
for i,row in enumerate(rows_t3):
    if i==0: header_row(t3.rows[0],row,w_t3)
    else:
        for j,(v,w) in enumerate(zip(row,w_t3)):
            t3.rows[i].cells[j].width=w
            is_lh=(j==0)
            cell_text(t3.rows[i].cells[j],v,bold=is_lh,center=(not is_lh),size=11)
            if is_lh: shade(t3.rows[i].cells[j],'D6EAF0')
            elif i%2==0: shade(t3.rows[i].cells[j],ALT)

spacer(doc)

# ════════════════════════════════════════════════════
#  TASK 3 — NORMATIVE SCALE
# ════════════════════════════════════════════════════
h2(doc,"Завдання 3. Побудова нормувальної шкали")

para(doc,
     "Нормувальна шкала дозволяє перетворити кількісний результат вимірювання "
     "на якісну оцінку рівня підготовленості відносно групи. Використовуємо "
     "5-рівневу сигмальну шкалу на основі M та σ.",indent=True)

h2(doc,"3.1 Формули меж шкали та їх розрахунок")

formula(doc,
        "Межа 1 / 2: x₁₂ = M − 1,5·σ")
formula(doc,
        f"x₁₂ = {Mb:.2f} − 1,5 · {sb:.2f} = {Mb:.2f} − {1.5*sb:.2f} = {Mb-1.5*sb:.2f} кгс")

formula(doc,"Межа 2 / 3: x₂₃ = M − 0,5·σ")
formula(doc,
        f"x₂₃ = {Mb:.2f} − 0,5 · {sb:.2f} = {Mb:.2f} − {0.5*sb:.2f} = {Mb-0.5*sb:.2f} кгс")

formula(doc,"Межа 3 / 4: x₃₄ = M + 0,5·σ")
formula(doc,
        f"x₃₄ = {Mb:.2f} + 0,5 · {sb:.2f} = {Mb:.2f} + {0.5*sb:.2f} = {Mb+0.5*sb:.2f} кгс")

formula(doc,"Межа 4 / 5: x₄₅ = M + 1,5·σ")
formula(doc,
        f"x₄₅ = {Mb:.2f} + 1,5 · {sb:.2f} = {Mb:.2f} + {1.5*sb:.2f} = {Mb+1.5*sb:.2f} кгс")

caption(doc,f"Таблиця 4. Нормувальна шкала  (M = {Mb:.2f} кгс,  σ = {sb:.2f} кгс)")

h_sc=["Рівень","Межі, кгс","Оцінка","Теоретична частка,%"]
w_sc=[Cm(4.5),Cm(5.0),Cm(1.8),Cm(3.5)]
scale_data=[
    h_sc,
    ["Низький",     f"менше {Mb-1.5*sb:.2f}",                          "1", "≈ 6,7%"],
    ["Нижче середнього",f"{Mb-1.5*sb:.2f} – {Mb-0.5*sb:.2f}",          "2","≈ 24,2%"],
    ["Середній",    f"{Mb-0.5*sb:.2f} – {Mb+0.5*sb:.2f}",              "3","≈ 38,2%"],
    ["Вище середнього",f"{Mb+0.5*sb:.2f} – {Mb+1.5*sb:.2f}",           "4","≈ 24,2%"],
    ["Високий",     f"більше {Mb+1.5*sb:.2f}",                          "5","≈ 6,7%"],
]
sc_bg=['','FADBD8','FAE5D3','D5F5E3','D6EAF8','F5EEF8']
t4=doc.add_table(rows=6,cols=4)
t4.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(t4)
for i,row in enumerate(scale_data):
    if i==0: header_row(t4.rows[0],row,w_sc)
    else:
        for j,(v,w) in enumerate(zip(row,w_sc)):
            t4.rows[i].cells[j].width=w
            cell_text(t4.rows[i].cells[j],v,center=True,size=12)
            shade(t4.rows[i].cells[j],sc_bg[i])

spacer(doc)

h2(doc,"3.2 Індивідуальна оцінка кожного спортсмена")

para(doc,
     "За отриманими межами шкали присвоюємо оцінку кожному спортсмену:",indent=True)

caption(doc,"Таблиця 5. Оцінка рівня силової підготовленості")

h_t5=["№","Спортсмен","Краща\nспроба, кгс","Порівняння з межами","Оцінка","Рівень"]
w_t5=[Cm(0.7),Cm(3.0),Cm(2.0),Cm(5.5),Cm(1.5),Cm(3.5)]
t5=doc.add_table(rows=n+1,cols=6)
t5.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(t5)
header_row(t5.rows[0],h_t5,w_t5)

score_bg_map={1:'FADBD8',2:'FAE5D3',3:'D5F5E3',4:'D6EAF8',5:'F5EEF8'}
b12=Mb-1.5*sb; b23=Mb-0.5*sb; b34=Mb+0.5*sb; b45=Mb+1.5*sb
for i in range(n):
    xv=grip_best[i]; sc=scores_b[i]
    if   sc==1: cmp_str=f"{xv:.0f} < {b12:.2f}"
    elif sc==2: cmp_str=f"{b12:.2f} ≤ {xv:.0f} < {b23:.2f}"
    elif sc==3: cmp_str=f"{b23:.2f} ≤ {xv:.0f} < {b34:.2f}"
    elif sc==4: cmp_str=f"{b34:.2f} ≤ {xv:.0f} < {b45:.2f}"
    else:       cmp_str=f"{xv:.0f} ≥ {b45:.2f}"
    row=t5.rows[i+1]
    bg=score_bg_map[sc]
    vals=[i+1,names[i],xv,cmp_str,sc,level_lbl[sc]]
    data_row(row,vals,w_t5,[True,False,True,False,True,True],bg=bg)

spacer(doc)

# distribution
h2(doc,"3.3 Розподіл групи за рівнями підготовленості")

from collections import Counter
cnt=Counter(scores_b)
para(doc,"Підрахуємо кількість спортсменів на кожному рівні:",indent=True)
for lv in [1,2,3,4,5]:
    c=cnt.get(lv,0)
    para(doc,
         f"  Рівень {lv} — «{level_lbl[lv]}»:  {c} спортсменів  "
         f"({c/n*100:.0f}%)",indent=False)

spacer(doc)

# ════════════════════════════════════════════════════
#  TASK 4 — DYNAMOGRAM
# ════════════════════════════════════════════════════
h2(doc,"Завдання 4. Аналіз динамограми. Розрахунок силових показників")

para(doc,
     f"Умова: записано динамограму відштовхування спортсмена від опори "
     f"(маса тіла m = {m_body_dyn:.0f} кг). Час відштовхування T = {T_tot*1000:.0f} мс. "
     f"Сила реакції опори F(t) зчитана у 15 рівновіддалених точках "
     f"(Δt = 25 мс = 0,025 с).",indent=True)

caption(doc,"Таблиця 6. Значення F(t) зчитані з динамограми")

h_dyn=["Точка i","t, мс","F(t), Н","Δt, с","(Fᵢ + Fᵢ₊₁)/2, Н","ΔI = (Fᵢ+Fᵢ₊₁)/2·Δt, Н·с"]
w_dyn=[Cm(1.4),Cm(1.5),Cm(2.0),Cm(1.5),Cm(3.0),Cm(3.0)]

# compute trapezoid contributions
trap_avg  = (F_dyn[:-1]+F_dyn[1:])/2   # 14 values
trap_dI   = trap_avg * dt_s             # 14 values

n_rows_dyn = len(t_ms)+2   # header + 15 data + sum
td=doc.add_table(rows=n_rows_dyn,cols=6)
td.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(td)
header_row(td.rows[0],h_dyn,w_dyn)

for i,(tm,fv) in enumerate(zip(t_ms,F_dyn)):
    row=td.rows[i+1]
    bg=ALT if i%2==0 else None
    if i < len(trap_avg):
        avg_str=f"{trap_avg[i]:.1f}"
        dI_str =f"{trap_dI[i]:.3f}"
    else:
        avg_str="—"; dI_str="—"
    vals=[i+1,f"{tm:.0f}",f"{fv:.0f}","0,025" if i<len(trap_avg) else "—",avg_str,dI_str]
    data_row(row,vals,w_dyn,[True,True,True,True,True,True],bg=bg)

# Sum row
sum_row=td.rows[n_rows_dyn-1]
for j,w in enumerate(w_dyn): sum_row.cells[j].width=w
for j,v in enumerate(["","","","Σ ΔI =","",f"{I_trap:.3f} Н·с"]):
    cell_text(sum_row.cells[j],v,bold=True,center=True)
    shade(sum_row.cells[j],'D6EAF0')

spacer(doc)

h2(doc,"4.1 Імпульс сили I")

para(doc,
     "Імпульс сили розраховуємо методом трапецій (чисельне інтегрування):",indent=True)

formula(doc,"I = ∫ F(t) dt  ≈  Σ (Fᵢ + Fᵢ₊₁)/2 · Δtᵢ")
formula(doc,
        f"I = {' + '.join(f'{v:.3f}' for v in trap_dI[:7])}")
formula(doc,
        f"    + {' + '.join(f'{v:.3f}' for v in trap_dI[7:])}")
formula(doc, f"I = {I_trap:.3f} ≈ {I_trap:.1f} Н·с")

spacer(doc)

h2(doc,"4.2 Середня сила F̄")

para(doc,
     "Середня сила — умовна постійна сила, яка за час T дає той самий імпульс:",indent=True)

formula(doc,"F̄ = I / T")
formula(doc,f"F̄ = {I_trap:.1f} / {T_tot:.3f}")
formula(doc,f"F̄ = {F_avg_I:.1f} Н")

spacer(doc)

h2(doc,"4.3 Максимальна сила F_max та час її досягнення")

para(doc,
     "З таблиці динамограми визначаємо:",indent=True)
formula(doc,f"F_max = {F_max_d:.0f} Н  (точка i = {np.argmax(F_dyn)+1},  t = {t_peak:.0f} мс)")

spacer(doc)

h2(doc,"4.4 Градієнт сили G_F")

para(doc,
     "Градієнт сили показує швидкість наростання сили від нуля до максимуму:",indent=True)

formula(doc,"G_F = F_max / t(F_max)")
formula(doc,f"G_F = {F_max_d:.0f} / {t_peak/1000:.3f}")
formula(doc,f"G_F = {G_F:.1f} Н/с")

spacer(doc)

h2(doc,"4.5 Відносна сила при відштовхуванні")

F_rel_push = F_max_d / F_body
formula(doc,"η = F_max / (m · g)  =  F_max / F_тіла")
formula(doc,f"η = {F_max_d:.0f} / (80 · 9,81) = {F_max_d:.0f} / {F_body:.1f}")
formula(doc,f"η = {F_rel_push:.2f}  (F_max перевищує вагу тіла у {F_rel_push:.2f} рази)")

caption(doc,"Таблиця 7. Зведені показники динамограми відштовхування")
dyn_summ=[
    ["Показник","Формула","Підстановка","Результат"],
    ["Максимальна сила F_max","F_max = max F(t)","max із таблиці",f"{F_max_d:.0f} Н"],
    ["Час до максимуму t(F_max)","—","з таблиці",f"{t_peak:.0f} мс = {t_peak/1000:.3f} с"],
    ["Загальний час T","—","задано","0,350 с"],
    ["Імпульс сили I","I ≈ Σ(Fᵢ+Fᵢ₊₁)/2·Δt",f"Σ із таблиці",f"{I_trap:.1f} Н·с"],
    ["Середня сила F̄","F̄ = I / T",f"{I_trap:.1f} / 0,350",f"{F_avg_I:.1f} Н"],
    ["Градієнт сили G_F","G_F = F_max / t(F_max)",f"{F_max_d:.0f} / {t_peak/1000:.3f}",f"{G_F:.1f} Н/с"],
    ["Відносна сила η","η = F_max / F_тіла",f"{F_max_d:.0f} / {F_body:.1f}",f"{F_rel_push:.2f}"],
]
w_ds=[Cm(4.0),Cm(4.0),Cm(3.0),Cm(2.5)]
tds=doc.add_table(rows=len(dyn_summ),cols=4)
tds.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(tds)
for i,row in enumerate(dyn_summ):
    if i==0: header_row(tds.rows[0],row,w_ds)
    else:
        for j,(v,w) in enumerate(zip(row,w_ds)):
            tds.rows[i].cells[j].width=w
            cell_text(tds.rows[i].cells[j],v,center=(j!=0),size=11)
            if i%2==0: shade(tds.rows[i].cells[j],ALT)

spacer(doc)

# ════════════════════════════════════════════════════
#  TASK 5 — RELATIVE STRENGTH
# ════════════════════════════════════════════════════
h2(doc,"Завдання 5. Відносна сила спортсменів")

para(doc,
     "Відносна сила — це відношення абсолютної сили до маси тіла. "
     "Вона дозволяє об'єктивно порівнювати спортсменів різної ваги.",indent=True)

formula(doc,"η = F_абс / m_тіла  (кгс/кг)")

para(doc,"Розраховуємо для кожного спортсмена:",indent=True)

for i in range(n):
    formula(doc,
            f"η_{i+1} = {grip_best[i]} / {body_weight[i]} = {rel_str[i]:.3f} кгс/кг  "
            f"({names[i]})")

spacer(doc)

# Normative scale for relative strength
para(doc,
     f"Статистичні характеристики відносної сили:",indent=True)
formula(doc,
        f"M_відн = ({' + '.join(f'{v:.3f}' for v in rel_str)}) / {n}")
formula(doc,f"M_відн = {np.sum(rel_str):.3f} / {n} = {Mr:.3f} кгс/кг")

dev_r = rel_str - Mr
sq_r  = dev_r**2
formula(doc,f"σ_відн = √(Σ(ηᵢ−M)² / (n−1)) = √({np.sum(sq_r):.6f} / {n-1})")
formula(doc,f"σ_відн = √{np.sum(sq_r)/(n-1):.6f} = {sr:.4f} кгс/кг")
formula(doc,f"V_відн = ({sr:.4f} / {Mr:.3f}) · 100% = {sr/Mr*100:.1f}%")

para(doc,
     "Нормувальна шкала для відносної сили:",bold=True,indent=False)
formula(doc,f"Рівень 1 (Низький):          η < {Mr-1.5*sr:.3f}")
formula(doc,f"Рівень 2 (Нижче середнього): {Mr-1.5*sr:.3f} ≤ η < {Mr-0.5*sr:.3f}")
formula(doc,f"Рівень 3 (Середній):         {Mr-0.5*sr:.3f} ≤ η < {Mr+0.5*sr:.3f}")
formula(doc,f"Рівень 4 (Вище середнього):  {Mr+0.5*sr:.3f} ≤ η < {Mr+1.5*sr:.3f}")
formula(doc,f"Рівень 5 (Високий):          η ≥ {Mr+1.5*sr:.3f}")

spacer(doc)

# Relative strength table
caption(doc,"Таблиця 8. Відносна сила та рейтинг спортсменів")
h_t8=["Місце","Спортсмен","Краща\nспроба, кгс","Маса\nтіла, кг","η = F/m\n(кгс/кг)","Формула","Оцінка\nабсол.","Оцінка\nвідн."]
w_t8=[Cm(1.2),Cm(2.8),Cm(1.9),Cm(1.5),Cm(2.0),Cm(2.8),Cm(1.7),Cm(1.7)]
idx_srt=np.argsort(grip_best)[::-1]
t8=doc.add_table(rows=n+1,cols=8)
t8.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(t8)
header_row(t8.rows[0],h_t8,w_t8)
medal_bg={1:GLD,2:SLV,3:BRZ}
for rank,(i) in enumerate(idx_srt,1):
    row=t8.rows[rank]
    bg=medal_bg.get(rank, ALT if rank%2==0 else None)
    form_str=f"{grip_best[i]}/{body_weight[i]}"
    vals=[rank,names[i],grip_best[i],body_weight[i],f"{rel_str[i]:.3f}",
          form_str,scores_b[i],scores_r[i]]
    data_row(row,vals,w_t8,[True,False,True,True,True,True,True,True],bg=bg)

spacer(doc)

# ════════════════════════════════════════════════════
#  TASK 6 — RELIABILITY
# ════════════════════════════════════════════════════
h2(doc,"Завдання 6. Визначення надійності тесту (r_tt)")

para(doc,
     "Надійність тесту — відтворюваність результатів при повторному вимірюванні. "
     "Оцінюємо за коефіцієнтом кореляції Пірсона між спробою 1 і спробою 2.",indent=True)

formula(doc,"r_tt = Σ(x₁ᵢ − M₁)·(x₂ᵢ − M₂)  /  √[ Σ(x₁ᵢ − M₁)² · Σ(x₂ᵢ − M₂)² ]")

para(doc,
     f"Середні значень: M₁ = {M1:.2f} кгс,  M₂ = {M2:.2f} кгс",indent=True)

caption(doc,"Таблиця 9. Розрахунок коефіцієнта кореляції Пірсона (r_tt)")

h_t9=["№","Спортсмен","x₁ᵢ","x₂ᵢ","d₁ = x₁−M₁","d₂ = x₂−M₂","d₁·d₂","d₁²","d₂²"]
w_t9=[Cm(0.7),Cm(2.8),Cm(1.2),Cm(1.2),Cm(1.6),Cm(1.6),Cm(1.6),Cm(1.5),Cm(1.5)]
t9=doc.add_table(rows=n+2,cols=9)
t9.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(t9)
header_row(t9.rows[0],h_t9,w_t9)

for i in range(n):
    row=t9.rows[i+1]
    bg=ALT if i%2==0 else None
    d1=grip_1[i]-M1; d2=grip_2[i]-M2
    vals=[i+1,names[i],grip_1[i],grip_2[i],
          f"{d1:.2f}",f"{d2:.2f}",f"{d1*d2:.2f}",f"{d1**2:.2f}",f"{d2**2:.2f}"]
    data_row(row,vals,w_t9,[True,False,True,True,True,True,True,True,True],bg=bg)

# sums row
sum_row9=t9.rows[n+1]
for j,w in enumerate(w_t9): sum_row9.cells[j].width=w
sums=["","Σ","","","",
      "",f"{num_r:.2f}",f"{np.sum(dev1**2):.2f}",f"{np.sum(dev2**2):.2f}"]
for j,v in enumerate(sums):
    cell_text(sum_row9.cells[j],v,bold=True,center=True)
    shade(sum_row9.cells[j],'D6EAF0')

spacer(doc)

h2(doc,"6.1 Підстановка у формулу")

formula(doc,f"Чисельник:   Σ d₁ᵢ·d₂ᵢ = {num_r:.2f}")
formula(doc,f"Знаменник:   √( Σd₁² · Σd₂² ) = √( {np.sum(dev1**2):.2f} · {np.sum(dev2**2):.2f} )")
formula(doc,f"             = √{np.sum(dev1**2)*np.sum(dev2**2):.2f} = {den_r:.4f}")
formula(doc,f"r_tt = {num_r:.2f} / {den_r:.4f}")
formula(doc,f"r_tt = {r_tt:.4f}")

spacer(doc)

h2(doc,"6.2 Перевірка значущості (t-критерій Стьюдента)")

formula(doc,"t = r_tt · √(n−2) / √(1 − r_tt²)")
formula(doc,
        f"t = {r_tt:.4f} · √({n}−2) / √(1 − {r_tt:.4f}²)")
formula(doc,
        f"t = {r_tt:.4f} · √{n-2} / √(1 − {r_tt**2:.6f})")
formula(doc,
        f"t = {r_tt:.4f} · {np.sqrt(n-2):.4f} / √{1-r_tt**2:.6f}")
formula(doc,
        f"t = {r_tt:.4f} · {np.sqrt(n-2):.4f} / {np.sqrt(1-r_tt**2):.6f}")
formula(doc, f"t_спост = {t_obs:.3f}")
formula(doc, f"t_кр (α = 0,05; df = {n}−2 = {n-2}) = {t_cr:.3f}")

if t_obs > t_cr:
    verdict = f"t_спост = {t_obs:.3f} > t_кр = {t_cr:.3f} → зв'язок СТАТИСТИЧНО ЗНАЧУЩИЙ (p < 0,001)"
else:
    verdict = f"t_спост = {t_obs:.3f} < t_кр = {t_cr:.3f} → зв'язок не значущий"
para(doc,verdict,bold=True,indent=True)

spacer(doc)

h2(doc,"6.3 Стандартна похибка вимірювання (SEM)")

formula(doc,"SEM = σ · √(1 − r_tt)")
formula(doc,f"SEM = {sb:.2f} · √(1 − {r_tt:.4f})")
formula(doc,f"SEM = {sb:.2f} · √{1-r_tt:.6f}")
formula(doc,f"SEM = {sb:.2f} · {np.sqrt(1-r_tt):.4f}")
formula(doc,f"SEM = {SEM:.3f} кгс")

para(doc,
     f"Це означає: при повторному вимірюванні результат того самого спортсмена "
     f"відхилиться від справжнього значення не більш ніж на ±{SEM:.2f} кгс (68% випадків).",indent=True)

caption(doc,"Таблиця 10. Критерії оцінки надійності та висновок")
rel_crit=[
    ["r_tt","Рівень надійності","Висновок щодо тесту"],
    ["≥ 0,90","Відмінна","Тест придатний для індивідуального контролю"],
    ["0,80–0,89","Хороша","Тест придатний для групового та індивідуального контролю"],
    ["0,70–0,79","Задовільна","Тест придатний лише для групового контролю"],
    ["< 0,70","Незадовільна","Тест ненадійний, потребує вдосконалення"],
]
w_rc=[Cm(2.0),Cm(3.5),Cm(8.0)]
trc=doc.add_table(rows=5,cols=3)
trc.alignment=WD_TABLE_ALIGNMENT.CENTER
set_borders(trc)
hl_map={'Відмінна':GRN,'Хороша':'D5F5E3','Задовільна':'FEF5E7','Незадовільна':RED}
for i,row in enumerate(rel_crit):
    if i==0: header_row(trc.rows[0],row,w_rc)
    else:
        for j,(v,w) in enumerate(zip(row,w_rc)):
            trc.rows[i].cells[j].width=w
            cell_text(trc.rows[i].cells[j],v,center=(j!=2),size=11)
            if rel_level in row[1]:
                shade(trc.rows[i].cells[j],GRN)
            elif i%2==0:
                shade(trc.rows[i].cells[j],ALT)

para(doc,
     f"Наш результат: r_tt = {r_tt:.4f} → рівень надійності «{rel_level}».",
     bold=True,indent=True)

spacer(doc)

# ════════════════════════════════════════════════════
#  TASK 7 — GRAPHS
# ════════════════════════════════════════════════════
h2(doc,"Завдання 7. Графічна частина")

caption(doc,"Графік 1. Динамограма відштовхування спортсмена від опори")
add_image(doc,'g1.png',15.5)
para(doc,
     f"На графіку: крива F(t) з 15 зафіксованими точками; заштрихована площа — "
     f"імпульс I = {I_trap:.1f} Н·с; горизонтальна пунктирна лінія — вага тіла "
     f"{F_body:.0f} Н; штрихпунктирна — середня сила F̄ = {F_avg_I:.0f} Н; "
     f"зелена стрілка показує градієнт G_F = {G_F:.0f} Н/с.",indent=True)

spacer(doc)

caption(doc,"Графік 2. Гістограма розподілу та нормувальна шкала")
add_image(doc,'g2.png',15.5)
para(doc,
     f"Ліва панель: гістограма та нормальна крива; вертикальні лінії — межі шкали. "
     f"Права панель: 5 кольорових рівнів, чорні риски — результати конкретних спортсменів.",
     indent=True)

spacer(doc)

caption(doc,"Графік 3. Кореляційне поле «спроба 1 – спроба 2»")
add_image(doc,'g3.png',13.0)
para(doc,
     f"Точки щільно розміщені вздовж лінії регресії (r_tt = {r_tt:.4f}), що "
     f"підтверджує відмінну надійність тесту.",indent=True)

spacer(doc)

caption(doc,"Графік 4. Порівняльний аналіз силової підготовленості")
add_image(doc,'g4.png',15.5)
para(doc,
     "Ліва панель — абсолютна сила (стовпці кольором відповідають рівню шкали); "
     "центральна — відносна сила; права — кругова діаграма розподілу за рівнями.",
     indent=True)

spacer(doc)

# ════════════════════════════════════════════════════
#  CONCLUSIONS
# ════════════════════════════════════════════════════
h1(doc,"ВИСНОВКИ")

top_i = np.argmax(grip_best)
low_i = np.argmin(grip_best)
top_r = np.argmax(rel_str)

concls=[
    ("1. Статистичний аналіз.",
     f"У групі n = {n} спортсменів середня максимальна сила кисті "
     f"M = {Mb:.2f} кгс, σ = {sb:.2f} кгс, V = {Vb:.1f}%, m = {mb:.2f} кгс. "
     f"Коефіцієнт варіації V = {Vb:.1f}% свідчить про "
     f"{'однорідну' if Vb<10 else 'помірно однорідну' if Vb<20 else 'неоднорідну'} "
     f"групу. Довірчий інтервал середнього: {Mb:.2f} ± {mb:.2f} кгс."),
    ("2. Нормувальна шкала.",
     f"На основі M = {Mb:.2f} та σ = {sb:.2f} кгс побудовано 5-рівневу шкалу. "
     f"Межі: {b12:.2f} / {b23:.2f} / {b34:.2f} / {b45:.2f} кгс. "
     f"Розподіл групи: Низький — {cnt.get(1,0)}/{n}; Нижче серед. — {cnt.get(2,0)}/{n}; "
     f"Середній — {cnt.get(3,0)}/{n}; Вище серед. — {cnt.get(4,0)}/{n}; "
     f"Високий — {cnt.get(5,0)}/{n}."),
    ("3. Динамограма.",
     f"Максимальна сила відштовхування F_max = {F_max_d:.0f} Н, "
     f"час досягнення — {t_peak:.0f} мс. "
     f"Імпульс сили I = {I_trap:.1f} Н·с (метод трапецій). "
     f"Середня сила F̄ = {F_avg_I:.0f} Н ({F_avg_I/F_max_d*100:.0f}% від F_max). "
     f"Градієнт сили G_F = {G_F:.0f} Н/с. "
     f"Відносна сила η = {F_rel_push:.2f} (F_max більша за вагу тіла у {F_rel_push:.2f} рази)."),
    ("4. Відносна сила.",
     f"Найвища абсолютна сила — {names[top_i]} ({grip_best[top_i]} кгс). "
     f"Найвища відносна сила — {names[top_r]} "
     f"({rel_str[top_r]:.3f} кгс/кг). "
     f"Найнижчий результат — {names[low_i]} ({grip_best[low_i]} кгс). "
     f"Порівняння абсолютної та відносної сили виявляє різницю в рейтингах "
     f"спортсменів, що важливо для видів спорту з ваговими категоріями."),
    ("5. Надійність тесту.",
     f"Коефіцієнт ретест-надійності r_tt = {r_tt:.4f} (рівень «{rel_level}»). "
     f"t_спост = {t_obs:.3f} > t_кр = {t_cr:.3f} → зв'язок статистично значущий (p < 0,001). "
     f"SEM = {SEM:.3f} кгс — дуже мала похибка вимірювання. "
     f"Тест динамометрії кисті є придатним і надійним інструментом "
     f"контролю максимальної сили."),
    ("6. Рекомендації.",
     f"Спортсменам з оцінками 1–2 ({cnt.get(1,0)+cnt.get(2,0)} осіб, "
     f"{(cnt.get(1,0)+cnt.get(2,0))/n*100:.0f}%) рекомендовано збільшити обсяг "
     f"силового тренування. При контролі слід враховувати відносну силу. "
     f"Для підвищення точності використовувати тензометричні пристрої "
     f"(r_tt = 0,85–0,95) замість механічних. "
     f"Обов'язково стандартизувати положення тіла при кожному вимірюванні."),
]

for bold_title, text in concls:
    p_c=doc.add_paragraph()
    p_c.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p_c.paragraph_format.first_line_indent=Cm(1.25)
    p_c.paragraph_format.space_after=Pt(6)
    r1=p_c.add_run(bold_title+" ")
    r1.font.name='Times New Roman'; r1.font.size=Pt(14); r1.font.bold=True
    r2=p_c.add_run(text)
    r2.font.name='Times New Roman'; r2.font.size=Pt(14)

spacer(doc)

# ════════════════════════════════════════════════════
#  Q&A FOR DEFENSE
# ════════════════════════════════════════════════════
h1(doc,"ТЕОРЕТИЧНІ ПИТАННЯ ДЛЯ ЗАХИСТУ")

qa=[
    ("1. Що таке силові якості та яке їх значення?",
     "Силові якості — здатність долати зовнішній опір за рахунок м'язових зусиль. "
     "Визначають результат у більшості видів спорту."),
    ("2. Які три групи показників використовують при контролі сили?",
     "1) Основні (F_max, F̄); 2) Інтегральні (імпульс I = ∫F dt); "
     "3) Диференціальні (градієнт G_F = dF/dt)."),
    ("3. Що таке імпульс сили і як його розраховують чисельно?",
     "I = ∫F(t) dt — площа під кривою F(t). Чисельно: метод трапецій "
     "I ≈ Σ (Fᵢ + Fᵢ₊₁)/2 · Δtᵢ."),
    ("4. Що таке градієнт сили і де він застосовується?",
     "G_F = F_max / t(F_max) — швидкість наростання сили. "
     "Важливий для стрибків, ударів, метань."),
    ("5. Як побудувати нормувальну шкалу?",
     "Розраховують M та σ вибірки. Межі шкали: M ± 0,5σ (рівні 2-3-4) "
     "та M ± 1,5σ (рівні 1 і 5)."),
    ("6. Що таке надійність тесту та як її оцінюють?",
     "Надійність — відтворюваність результату. Оцінюється коефіцієнтом "
     "кореляції r_tt між повторними вимірюваннями. "
     "r_tt ≥ 0,90 — відмінна; 0,80–0,89 — хороша; 0,70–0,79 — задовільна."),
    ("7. Що таке відносна сила і чому вона важлива?",
     "η = F_абс / m_тіла (кгс/кг). Дозволяє об'єктивно порівнювати "
     "спортсменів різної маси. Обов'язкова в єдиноборствах і гімнастиці."),
    ("8. Яка різниця між статичним і динамічним вимірюванням сили?",
     "Статичне — ізометрично (без руху); динамічне — в русі. "
     "Інерційні динамографи дозволяють вимірювати силу безпосередньо в русі, "
     "що інформативніше для більшості видів спорту."),
]

for q,a in qa:
    pq=doc.add_paragraph()
    pq.paragraph_format.space_before=Pt(8); pq.paragraph_format.space_after=Pt(2)
    rq=pq.add_run(q); rq.font.name='Times New Roman'; rq.font.size=Pt(14); rq.font.bold=True
    pa=doc.add_paragraph()
    pa.paragraph_format.left_indent=Cm(1.0); pa.paragraph_format.space_after=Pt(4)
    ra=pa.add_run("Відповідь: "+a)
    ra.font.name='Times New Roman'; ra.font.size=Pt(14)

spacer(doc)

# ════════════════════════════════════════════════════
#  BIBLIOGRAPHY
# ════════════════════════════════════════════════════
h1(doc,"СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ")

refs=[
    "Зациорский В.М. Спортивная метрология. — М.: Физкультура и спорт, 1982. — 256 с.",
    "Платонов В.Н. Система подготовки спортсменов в олимпийском спорте. — К.: Олимпийская "
    "литература, 2004. — 808 с.",
    "Годик М.А. Спортивная метрология: учебник для ИФК. — М.: Физкультура и спорт, 1988. — 192 с.",
    "Круцевич Т.Ю. Теорія і методика фізичного виховання. — К.: Олімпійська "
    "література, 2008. — Т. 1. — 392 с.",
    "Сергієнко Л.П. Спортивна метрологія: теорія і практичні аспекти. — К.: КНТ, 2010. — 776 с.",
]
for i,ref in enumerate(refs,1):
    p_r=doc.add_paragraph()
    p_r.paragraph_format.left_indent=Cm(1.25)
    p_r.paragraph_format.first_line_indent=Cm(-1.25)
    p_r.paragraph_format.space_after=Pt(4)
    r_r=p_r.add_run(f"{i}. {ref}")
    r_r.font.name='Times New Roman'; r_r.font.size=Pt(13)

# ─────────────────────────────────────────────────
out='practical_work_18_full.docx'
doc.save(out)
print(f"\nДокумент збережено: {out}")
print(f"Розмір: {os.path.getsize(out)//1024} КБ")
