"""
Generates practical_work_18.docx — Practical Work 18
"Metrological foundations of control over the level of development of physical strength quality"
"""

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from scipy import stats
from scipy.integrate import trapezoid
import warnings
warnings.filterwarnings('ignore')

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy
import os

# ============================================================
# ВИХІДНІ ДАНІ
# ============================================================
names = [
    'Андрієнко', 'Бойко', 'Василенко', 'Гнатенко', 'Демченко',
    'Єременко', 'Захаренко', 'Іваненко', 'Коваленко', 'Литвиненко',
    'Марченко', 'Назаренко', 'Олексієнко', 'Петренко', 'Романенко',
    'Савченко', 'Тимченко', 'Удовенко', 'Філіпенко', 'Харченко'
]
body_weight = np.array([75,82,68,90,78,85,72,88,76,81,70,93,65,87,79,84,73,91,77,83])
grip_1 = np.array([48,55,42,61,51,58,46,63,49,54,44,67,40,59,52,57,47,65,50,56])
grip_2 = np.array([47,56,43,60,52,59,45,62,50,55,45,66,41,60,53,58,48,64,51,55])
grip_best = np.maximum(grip_1, grip_2)
relative_strength = grip_best / body_weight

n_s  = len(grip_best)
M    = np.mean(grip_best)
sig  = np.std(grip_best, ddof=1)

def assign_score(x, mean, sd):
    if   x < mean - 1.5*sd: return 1
    elif x < mean - 0.5*sd: return 2
    elif x < mean + 0.5*sd: return 3
    elif x < mean + 1.5*sd: return 4
    else:                    return 5

level_map = {1:'Низький',2:'Нижче середнього',3:'Середній',4:'Вище середнього',5:'Високий'}
scores_list   = [assign_score(v, M, sig)       for v in grip_best]
rel_str_scores = [assign_score(v, np.mean(relative_strength), np.std(relative_strength,ddof=1))
                  for v in relative_strength]

r_tt, p_value = stats.pearsonr(grip_1, grip_2)
t_obs = r_tt * np.sqrt(n_s-2) / np.sqrt(1-r_tt**2)
t_cr  = stats.t.ppf(0.975, df=n_s-2)
SEM_meas = np.std(grip_best, ddof=1) * np.sqrt(1 - r_tt)
rel_level = 'ВІДМІННА' if r_tt >= 0.90 else ('ХОРОША' if r_tt >= 0.80 else 'ЗАДОВІЛЬНА')

# Dynamogram
T_total = 0.35; dt = 0.001
t_arr   = np.arange(0, T_total+dt, dt)
t_peak  = 0.12; F_max_dyn = 2200; m_body = 80; F_body = m_body*9.81
F_arr   = np.where(
    t_arr <= t_peak,
    F_body + (F_max_dyn-F_body)*np.sin(np.pi/2*t_arr/t_peak)**1.6,
    F_body + (F_max_dyn-F_body)*np.exp(-4.5*(t_arr-t_peak)/(T_total-t_peak))
)
np.random.seed(42)
F_arr   += np.random.normal(0,15,size=len(t_arr))
F_arr   = np.clip(F_arr,0,None)
F_max_val   = np.max(F_arr)
t_Fmax      = t_arr[np.argmax(F_arr)]
impulse     = trapezoid(F_arr, t_arr)
F_avg_from_I= impulse/T_total
gradient_F  = F_max_val/t_Fmax


# ============================================================
# GENERATE GRAPHS
# ============================================================
plt.rcParams['font.family'] = 'DejaVu Sans'
plt.rcParams['axes.grid']   = True
plt.rcParams['grid.alpha']  = 0.3

def save_graph1():
    fig, ax = plt.subplots(figsize=(13,6))
    ax.plot(t_arr*1000, F_arr, color='royalblue', lw=2, label='F(t) – сила реакції опори')
    ax.axhline(F_body, color='gray', ls='--', lw=1.2, label=f'Вага тіла = {F_body:.0f} Н')
    ax.axhline(F_avg_from_I, color='orange', ls='-.', lw=1.5,
               label=f'Середня сила = {F_avg_from_I:.0f} Н')
    ax.fill_between(t_arr*1000, F_arr, alpha=0.12, color='royalblue',
                    label=f'Імпульс I = {impulse:.1f} Н·с')
    ax.annotate(
        f'F_max = {F_max_val:.0f} Н\nt = {t_Fmax*1000:.0f} мс',
        xy=(t_Fmax*1000, F_max_val), xytext=(t_Fmax*1000+30, F_max_val-200),
        arrowprops=dict(arrowstyle='->', color='red', lw=1.8),
        fontsize=11, color='red',
        bbox=dict(boxstyle='round,pad=0.3', fc='lightyellow', ec='red')
    )
    ax.annotate('', xy=(t_Fmax*1000,F_max_val), xytext=(0,F_body),
                arrowprops=dict(arrowstyle='->', color='green', lw=1.5, linestyle='dashed'))
    ax.text(t_Fmax*500-15,(F_max_val+F_body)/2,
            f'G_F = {gradient_F:.0f} Н/с', color='green', fontsize=10)
    ax.set_xlabel('Час, мс', fontsize=13)
    ax.set_ylabel('Сила, Н', fontsize=13)
    ax.set_title('Графік 1. Динамограма відштовхування спортсмена від опори', fontsize=13, fontweight='bold')
    ax.legend(loc='upper right', fontsize=10)
    ax.set_xlim(0, T_total*1000); ax.set_ylim(0, F_max_val*1.15)
    plt.tight_layout()
    plt.savefig('g1.png', dpi=150, bbox_inches='tight'); plt.close()

def save_graph2():
    fig, (ax1,ax2) = plt.subplots(1,2,figsize=(14,6))
    bins = np.arange(38,70,4)
    counts,_,_ = ax1.hist(grip_best, bins=bins, edgecolor='white', lw=1.5,
                           color='steelblue', alpha=0.75)
    x_r = np.linspace(grip_best.min()-5, grip_best.max()+5, 200)
    ax1.plot(x_r, stats.norm.pdf(x_r,M,sig)*n_s*(bins[1]-bins[0]),
             color='navy', lw=2.5, label='Нормальний розподіл')
    scale_cols = ['#ef5350','#ff9800','#66bb6a','#42a5f5','#ab47bc']
    bounds_v   = [M-1.5*sig, M-0.5*sig, M, M+0.5*sig, M+1.5*sig]
    labels_v   = ['M-1.5σ','M-0.5σ','M','M+0.5σ','M+1.5σ']
    for b,col,lbl in zip(bounds_v, scale_cols, labels_v):
        ax1.axvline(b, color=col, ls='--', lw=1.8, alpha=0.85)
        ax1.text(b, counts.max()*1.05, lbl, ha='center', fontsize=9, color=col)
    ax1.set_xlabel('Сила кисті, кгс', fontsize=12)
    ax1.set_ylabel('Кількість спортсменів', fontsize=12)
    ax1.set_title('Розподіл результатів та межі нормувальної шкали', fontsize=11, fontweight='bold')
    ax1.legend(fontsize=10)

    level_names = ['Низький','Нижче\nсередн.','Середній','Вище\nсередн.','Високий']
    lo_b = [0, M-1.5*sig, M-0.5*sig, M+0.5*sig, M+1.5*sig]
    hi_b = [M-1.5*sig, M-0.5*sig, M+0.5*sig, M+1.5*sig, 80]
    for i,(lo,hi,col,nm) in enumerate(zip(lo_b,hi_b,scale_cols,level_names)):
        ax2.barh(i,hi-lo,left=lo,color=col,alpha=0.75,edgecolor='white',height=0.7)
        ax2.text((lo+hi)/2,i,nm,ha='center',va='center',fontsize=11,fontweight='bold',color='white')
        ax2.text(lo+0.3,i-0.42,f'{lo:.1f}',ha='left',fontsize=9)
        if hi<80: ax2.text(hi-0.3,i-0.42,f'{hi:.1f}',ha='right',fontsize=9)
    for val in grip_best:
        lvl = assign_score(val,M,sig)-1
        ax2.plot(val,lvl,'k|',ms=16,mew=2,alpha=0.6)
    ax2.set_xlabel('Сила кисті, кгс', fontsize=12)
    ax2.set_yticks(range(5))
    ax2.set_yticklabels(['1–Низький','2–Нижче серед.','3–Середній','4–Вище серед.','5–Високий'],fontsize=10)
    ax2.set_title('Нормувальна шкала (риски – значення спортсменів)', fontsize=11, fontweight='bold')
    ax2.set_xlim(30,75)
    plt.tight_layout()
    plt.savefig('g2.png', dpi=150, bbox_inches='tight'); plt.close()

def save_graph3():
    fig, ax = plt.subplots(figsize=(8,8))
    ax.scatter(grip_1, grip_2, s=90, color='royalblue', edgecolors='navy', alpha=0.8, zorder=5)
    slope,intercept,*_ = stats.linregress(grip_1, grip_2)
    x_lr = np.linspace(grip_1.min()-2, grip_1.max()+2, 100)
    ax.plot(x_lr, slope*x_lr+intercept, color='red', lw=2,
            label=f'Регресія: y = {slope:.2f}x + {intercept:.2f}')
    lim_mn = min(grip_1.min(),grip_2.min())-2
    lim_mx = max(grip_1.max(),grip_2.max())+2
    ax.plot([lim_mn,lim_mx],[lim_mn,lim_mx],'g--',lw=1.5,alpha=0.5,label='Бісектриса')
    for i,nm in enumerate(names):
        ax.annotate(nm[:4],(grip_1[i],grip_2[i]),xytext=(3,3),
                    textcoords='offset points',fontsize=7,color='gray')
    ax.set_xlabel('Спроба 1 (кгс)', fontsize=13)
    ax.set_ylabel('Спроба 2 (кгс)', fontsize=13)
    ax.set_title(f'Графік 3. Кореляційне поле (надійність тесту)\nr_tt = {r_tt:.4f}  (p < 0.001)',
                 fontsize=12, fontweight='bold')
    ax.legend(fontsize=10)
    ax.set_xlim(lim_mn,lim_mx); ax.set_ylim(lim_mn,lim_mx); ax.set_aspect('equal')
    plt.tight_layout()
    plt.savefig('g3.png', dpi=150, bbox_inches='tight'); plt.close()

def save_graph4():
    fig, axes = plt.subplots(1,3,figsize=(18,6))
    cols_bar  = [plt.cm.RdYlGn(assign_score(v,M,sig)/5) for v in grip_best]
    axes[0].bar(range(1,n_s+1), grip_best, color=cols_bar, edgecolor='white', lw=0.8)
    axes[0].axhline(M, color='navy', ls='--', lw=2, label=f'M = {M:.1f}')
    axes[0].axhline(M+sig, color='green', ls=':', lw=1.5, label=f'M+σ = {M+sig:.1f}')
    axes[0].axhline(M-sig, color='orange', ls=':', lw=1.5, label=f'M-σ = {M-sig:.1f}')
    axes[0].set_xticks(range(1,n_s+1))
    axes[0].set_xticklabels([n[:4] for n in names],rotation=45,ha='right',fontsize=8)
    axes[0].set_ylabel('Сила кисті, кгс', fontsize=11)
    axes[0].set_title('Абсолютна сила кисті', fontsize=12, fontweight='bold')
    axes[0].legend(fontsize=9)

    M_rel=np.mean(relative_strength); sd_rel=np.std(relative_strength,ddof=1)
    cols_rel=[plt.cm.RdYlGn(assign_score(v,M_rel,sd_rel)/5) for v in relative_strength]
    axes[1].bar(range(1,n_s+1), relative_strength, color=cols_rel, edgecolor='white', lw=0.8)
    axes[1].axhline(M_rel, color='navy', ls='--', lw=2, label=f'M = {M_rel:.3f}')
    axes[1].set_xticks(range(1,n_s+1))
    axes[1].set_xticklabels([n[:4] for n in names],rotation=45,ha='right',fontsize=8)
    axes[1].set_ylabel('Відносна сила, кгс/кг', fontsize=11)
    axes[1].set_title('Відносна сила (кгс/кг тіла)', fontsize=12, fontweight='bold')
    axes[1].legend(fontsize=9)

    level_order=['Низький','Нижче середнього','Середній','Вище середнього','Високий']
    level_cols_p=['#ef5350','#ff9800','#66bb6a','#42a5f5','#ab47bc']
    from collections import Counter
    cnt = Counter([level_map[s] for s in scores_list])
    sizes=[cnt.get(l,0) for l in level_order]
    non_zero=[(s,l,c) for s,l,c in zip(sizes,level_order,level_cols_p) if s>0]
    s_nz,l_nz,c_nz=zip(*non_zero)
    axes[2].pie(s_nz,labels=l_nz,colors=c_nz,
                autopct=lambda p:f'{p:.0f}%\n({int(round(p*n_s/100))} ос.)',
                startangle=90,pctdistance=0.75,
                wedgeprops=dict(edgecolor='white',linewidth=2))
    for t_ in axes[2].texts: t_.set_fontsize(9)
    axes[2].set_title('Розподіл за рівнями\nсилової підготовленості', fontsize=12, fontweight='bold')
    plt.suptitle('Графік 4. Порівняльний аналіз силової підготовленості',
                 fontsize=13,fontweight='bold',y=1.02)
    plt.tight_layout()
    plt.savefig('g4.png', dpi=150, bbox_inches='tight'); plt.close()

print("Generating graphs…")
save_graph1(); save_graph2(); save_graph3(); save_graph4()
print("Graphs saved.")


# ============================================================
# HELPER FUNCTIONS FOR DOCX
# ============================================================

def set_doc_defaults(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.paragraph_format.space_after = Pt(0)

def set_margins(doc, top=2, bottom=2, left=3, right=1.5):
    for section in doc.sections:
        section.top_margin    = Cm(top)
        section.bottom_margin = Cm(bottom)
        section.left_margin   = Cm(left)
        section.right_margin  = Cm(right)

def heading(doc, text, level=1):
    p = doc.add_paragraph(text)
    run = p.runs[0]
    run.font.name = 'Times New Roman'
    run.font.bold = True
    run.font.size = Pt(14) if level > 1 else Pt(16)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    return p

def body(doc, text, indent=False, bold_parts=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    return p

def bold_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after  = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.italic = True
    return p

def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, center=False, size=12, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(str(text))
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def set_table_borders(table):
    tbl  = table._tbl
    tblPr= tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    tblBorders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'),   'single')
        tag.set(qn('w:sz'),    '4')
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), '000000')
        tblBorders.append(tag)
    tblPr.append(tblBorders)

def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.font.name  = 'Times New Roman'
    run.font.size  = Pt(13)
    run.font.italic= True

def add_image(doc, path, width_cm=16):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Cm(width_cm))

def section_break(doc):
    doc.add_paragraph()


# ============================================================
# BUILD DOCUMENT
# ============================================================

doc = Document()
set_doc_defaults(doc)
set_margins(doc)

# ── TITLE PAGE ─────────────────────────────────────────────

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("ПРАКТИЧНА РОБОТА № 18")
run.font.name = 'Times New Roman'; run.font.size = Pt(18); run.font.bold = True

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = p2.add_run("Метрологічні основи контролю за рівнем розвитку\nфізичної якості сили")
run2.font.name = 'Times New Roman'; run2.font.size = Pt(16); run2.font.bold = True
p2.paragraph_format.space_after = Pt(8)

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
run3 = p3.add_run("Мета заняття: навчитися методиці оцінки силових якостей спортсменів.")
run3.font.name = 'Times New Roman'; run3.font.size = Pt(14); run3.font.italic = True
p3.paragraph_format.space_after = Pt(18)

doc.add_paragraph()

# ── SECTION 1 ──────────────────────────────────────────────

heading(doc, "1. ВИЗНАЧЕННЯ СИЛИ ЯК ФІЗИЧНОЇ ЯКОСТІ ЛЮДИНИ", level=1)

body(doc,
     "Здатність долати зовнішній опір або протидіяти йому за допомогою м'язових зусиль "
     "називають силовими якостями. Від рівня їх розвитку залежать досягнення фактично "
     "у всіх видах спорту, і тому методам контролю і вдосконалення силових якостей "
     "надається значна увага. Методи контролю за силовими якостями мають давню "
     "історію. Перші механічні пристрої для вимірювання сили були створені ще у XVIII ст.",
     indent=True)

heading(doc, "1.1 Групи показників при контролі силових якостей", level=2)

body(doc, "При контролі за силовими якостями враховують три групи показників:", indent=True)

groups = [
    ("1. Основні:", "а) миттєві значення сили у будь-який момент руху, зокрема максимальна сила; "
                    "б) середня сила."),
    ("2. Інтегральні:", "імпульс сили — площа під кривою F(t)."),
    ("3. Диференціальні:", "градієнт сили — показує, як швидко змінюються миттєві величини сили."),
]
for bold_part, rest in groups:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent   = Cm(1.25)
    p.paragraph_format.space_after   = Pt(4)
    r1 = p.add_run(bold_part + " ")
    r1.font.name='Times New Roman'; r1.font.size=Pt(14); r1.font.bold=True
    r2 = p.add_run(rest)
    r2.font.name='Times New Roman'; r2.font.size=Pt(14)

heading(doc, "1.2 Фізичні основи показників", level=2)

body(doc,
     "Згідно із законами механіки кінцевий ефект дії сили визначається імпульсом сили "
     "(позначається I). Графічно — це площа, обмежена кривою F(t).",
     indent=True)

add_formula(doc, "I = ∫ F(t) dt     або, при постійній силі:     I = F · t")

body(doc,
     "Середня сила — умовний показник, рівний частці від ділення імпульсу на час дії сили:",
     indent=True)

add_formula(doc, "F̄ = I / t = ∫F(t) dt / t")

body(doc,
     "Градієнт сили — диференціальний показник, що характеризує швидкість наростання сили:",
     indent=True)

add_formula(doc, "G_F = dF/dt      або практично:      G_F = F_max / t(F_max)")

section_break(doc)

# ── SECTION 2 ──────────────────────────────────────────────

heading(doc, "2. ВИДИ СИЛИ", level=1)

body(doc,
     "Максимальна сила — найвищий рівень напруги, що може розвинути м'яз. Наочна, але "
     "у швидких рухах порівняно погано характеризує їх кінцевий результат (кореляція між "
     "максимальною силою відштовхування та висотою стрибка може бути близька до нуля).",
     indent=True)

body(doc,
     "При метанні предметів різної маси виділяють три зони: зона I (малі обтяження) — "
     "характеризує швидкісні якості; зона II (середні) — швидкісно-силові; "
     "зона III (значні обтяження) — силові якості.",
     indent=True)

# Таблиця видів сили
heading(doc, "Таблиця 1. Класифікація видів сили", level=2)

tbl_kinds_data = [
    ["Вид сили", "Характеристика", "Застосування"],
    ["Максимальна сила", "Найвищий рівень напруги м'яза без урахування часу", "Важка атлетика, боротьба"],
    ["Вибухова сила", "Здатність досягти максимуму сили за мінімальний час", "Стрибки, метання, спринт"],
    ["Силова витривалість", "Тривале підтримання значних силових зусиль", "Веслування, плавання"],
    ["Відносна сила", "Максимальна сила / маса тіла спортсмена", "Єдиноборства (вагові категорії)"],
    ["Динамічна сила", "Сила, що проявляється в русі", "Більшість видів спорту"],
    ["Статична сила", "Ізометрична сила (без руху)", "Важка атлетика, боротьба"],
]

tbl_k = doc.add_table(rows=len(tbl_kinds_data), cols=3)
tbl_k.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_k)
col_widths_k = [Cm(4.5), Cm(7.5), Cm(5.5)]
for i, row_data in enumerate(tbl_kinds_data):
    row = tbl_k.rows[i]
    for j, val in enumerate(row_data):
        row.cells[j].width = col_widths_k[j]
        is_header = (i == 0)
        set_cell_text(row.cells[j], val, bold=is_header, center=is_header, size=12)
        if is_header:
            shade_cell(row.cells[j], '1F5C8B')
            run_ = row.cells[j].paragraphs[0].runs[0]
            run_.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        elif i % 2 == 0:
            shade_cell(row.cells[j], 'EAF2FB')

section_break(doc)

# ── SECTION 3 ──────────────────────────────────────────────

heading(doc, "3. КОНТРОЛЬ ЗА РІВНЕМ РОЗВИТКУ СИЛИ", level=1)

heading(doc, "3.1 Способи реєстрації силових якостей", level=2)

for txt in [
    "Розрізняють два способи реєстрації силових якостей:",
    "1) Без вимірювальної апаратури — оцінка за найбільшою вагою, яку здатний підняти або "
    "утримати спортсмен (прямий спосіб: жим штанги лежачи; непрямий: стрибки, метання).",
    "2) З використанням вимірювальних пристроїв — динамометрів або динамографів.",
]:
    body(doc, txt, indent=True)

heading(doc, "3.2 Типи силовимірювальних установок", level=2)

body(doc,
     "Всі силовимірювальні установки поділяються на два типи: а) що вимірюють деформацію "
     "тіла, до якого прикладена сила; б) інерційні динамографи — вимірюють прискорення "
     "рухомого тіла (перевага — вимірювання сили в русі, а не в статиці).",
     indent=True)

heading(doc, "3.3 Надійність силових тестів", level=2)

# Таблиця надійності
tbl_rel_data = [
    ["Тип тесту / пристрою", "Коефіцієнт надійності r_tt"],
    ["Механічні динамометри (пружинні)", "0,60 – 0,80"],
    ["Градієнти сили (будь-який спосіб)", "0,70 – 0,80"],
    ["Тензометричні пристрої (максимальна сила)", "0,85 – 0,95"],
]
tbl_rel = doc.add_table(rows=len(tbl_rel_data), cols=2)
tbl_rel.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_rel)
for i, row_data in enumerate(tbl_rel_data):
    row = tbl_rel.rows[i]
    row.cells[0].width = Cm(10); row.cells[1].width = Cm(6)
    is_h = (i == 0)
    set_cell_text(row.cells[0], row_data[0], bold=is_h, center=is_h, size=12)
    set_cell_text(row.cells[1], row_data[1], bold=is_h, center=True, size=12)
    if is_h:
        for c in row.cells:
            shade_cell(c, '1F5C8B')
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    elif i % 2 == 0:
        for c in row.cells: shade_cell(c, 'EAF2FB')

section_break(doc)

# ── RGR-2 TITLE ────────────────────────────────────────────

heading(doc, "РОЗРАХУНКОВО-ГРАФІЧНА РОБОТА № 2", level=1)
heading(doc, "«Метрологічний аналіз показників силової підготовленості спортсменів»", level=1)

body(doc, "Завдання роботи:", indent=False)
for task_txt in [
    "1. Розрахувати основні статистичні показники вибірки результатів вимірювання сили.",
    "2. Побудувати нормувальну шкалу для оцінки рівня силової підготовленості.",
    "3. Розрахувати імпульс сили, середню силу та градієнт сили за динамограмою.",
    "4. Розрахувати відносну силу спортсменів.",
    "5. Визначити надійність тесту (коефіцієнт кореляції між двома спробами).",
    "6. Побудувати графіки та зробити висновки.",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(task_txt)
    run.font.name='Times New Roman'; run.font.size=Pt(14)

section_break(doc)

# ── ЗАВДАННЯ 1: ПЕРВИННІ ДАНІ ───────────────────────────────

heading(doc, "Завдання 1. Первинні дані та статистичний аналіз", level=2)

body(doc,
     "Проведено динамометрію кисті (вимірювання максимальної сили стискання кисті) у групі "
     "спортсменів (n = 20). Вимірювання проведено двічі (спроба 1 і спроба 2) для оцінки "
     "надійності тесту. Одиниця вимірювання — кгс (кілограм-сила).",
     indent=True)

caption(doc, "Таблиця 2. Результати динамометрії кисті (n = 20 спортсменів)")

headers_t2 = ["№", "Спортсмен", "Маса тіла\n(кг)", "Спроба 1\n(кгс)", "Спроба 2\n(кгс)",
              "Краща\nспроба (кгс)", "Відносна\nсила (кгс/кг)"]
col_w_t2   = [Cm(0.9), Cm(3.5), Cm(2.0), Cm(2.0), Cm(2.0), Cm(2.3), Cm(2.5)]

tbl2 = doc.add_table(rows=n_s+1, cols=7)
tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl2)

for j,(h,w) in enumerate(zip(headers_t2, col_w_t2)):
    tbl2.rows[0].cells[j].width = w
    set_cell_text(tbl2.rows[0].cells[j], h, bold=True, center=True, size=11)
    shade_cell(tbl2.rows[0].cells[j], '1F5C8B')
    tbl2.rows[0].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

for i in range(n_s):
    row_d = tbl2.rows[i+1]
    vals  = [i+1, names[i], body_weight[i], grip_1[i], grip_2[i],
             grip_best[i], f"{relative_strength[i]:.3f}"]
    for j,(val,w) in enumerate(zip(vals,col_w_t2)):
        row_d.cells[j].width = w
        set_cell_text(row_d.cells[j], val, center=(j!=1), size=11)
        if i % 2 == 0:
            shade_cell(row_d.cells[j], 'F5F5F5')

section_break(doc)

# Статистика
heading(doc, "1.1 Описова статистика", level=2)

stats_rows = [
    ["Показник", "Спроба 1", "Спроба 2", "Краща спроба", "Відносна сила"],
    ["M (середнє)",
     f"{np.mean(grip_1):.2f} кгс", f"{np.mean(grip_2):.2f} кгс",
     f"{M:.2f} кгс", f"{np.mean(relative_strength):.3f} кгс/кг"],
    ["σ (відхилення)",
     f"{np.std(grip_1,ddof=1):.2f}", f"{np.std(grip_2,ddof=1):.2f}",
     f"{sig:.2f}", f"{np.std(relative_strength,ddof=1):.3f}"],
    ["V% (варіація)",
     f"{np.std(grip_1,ddof=1)/np.mean(grip_1)*100:.1f}%",
     f"{np.std(grip_2,ddof=1)/np.mean(grip_2)*100:.1f}%",
     f"{sig/M*100:.1f}%", f"{np.std(relative_strength,ddof=1)/np.mean(relative_strength)*100:.1f}%"],
    ["m (похибка середнього)",
     f"{np.std(grip_1,ddof=1)/np.sqrt(n_s):.2f}", f"{np.std(grip_2,ddof=1)/np.sqrt(n_s):.2f}",
     f"{sig/np.sqrt(n_s):.2f}", f"{np.std(relative_strength,ddof=1)/np.sqrt(n_s):.4f}"],
    ["Min",
     f"{np.min(grip_1):.0f}", f"{np.min(grip_2):.0f}",
     f"{np.min(grip_best):.0f}", f"{np.min(relative_strength):.3f}"],
    ["Max",
     f"{np.max(grip_1):.0f}", f"{np.max(grip_2):.0f}",
     f"{np.max(grip_best):.0f}", f"{np.max(relative_strength):.3f}"],
    ["Медіана",
     f"{np.median(grip_1):.1f}", f"{np.median(grip_2):.1f}",
     f"{np.median(grip_best):.1f}", f"{np.median(relative_strength):.3f}"],
]

caption(doc, "Таблиця 3. Описова статистика показників силової підготовленості")
tbl3 = doc.add_table(rows=len(stats_rows), cols=5)
tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl3)
cw3 = [Cm(5.5), Cm(2.7), Cm(2.7), Cm(2.7), Cm(2.7)]
for i, row_d in enumerate(stats_rows):
    for j, (val, w) in enumerate(zip(row_d, cw3)):
        tbl3.rows[i].cells[j].width = w
        is_h = (i == 0) or (j == 0)
        set_cell_text(tbl3.rows[i].cells[j], val, bold=is_h, center=(j != 0 or i == 0), size=11)
        if i == 0:
            shade_cell(tbl3.rows[i].cells[j], '1F5C8B')
            tbl3.rows[i].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        elif j == 0:
            shade_cell(tbl3.rows[i].cells[j], 'D6E4F0')
        elif i % 2 == 0:
            shade_cell(tbl3.rows[i].cells[j], 'F5F5F5')

body(doc,
     f"Висновок: коефіцієнт варіації кращої спроби V = {sig/M*100:.1f}% — група "
     f"{'однорідна (V < 10%)' if sig/M*100 < 10 else 'помірно однорідна (10% ≤ V < 20%)' if sig/M*100 < 20 else 'неоднорідна (V ≥ 20%)'}.",
     indent=True)

section_break(doc)

# ── ЗАВДАННЯ 2: НОРМУВАЛЬНА ШКАЛА ──────────────────────────

heading(doc, "Завдання 2. Нормувальна шкала оцінки силової підготовленості", level=2)

body(doc,
     "На основі середнього (M) та стандартного відхилення (σ) будується п'ятирівнева "
     "сигмальна нормувальна шкала з межами M ± 0,5σ та M ± 1,5σ.",
     indent=True)

scale_rows = [
    ["Рівень", "Нижня межа, кгс", "Верхня межа, кгс", "Оцінка", "% спортсменів"],
    ["Низький",            "< 42,7",  "< 42,7",  "1",
     f"{sum(1 for s in scores_list if s==1)/n_s*100:.0f}%"],
    ["Нижче середнього",   "42,7",    "50,1",     "2",
     f"{sum(1 for s in scores_list if s==2)/n_s*100:.0f}%"],
    ["Середній",           "50,1",    "57,6",     "3",
     f"{sum(1 for s in scores_list if s==3)/n_s*100:.0f}%"],
    ["Вище середнього",    "57,6",    "65,0",     "4",
     f"{sum(1 for s in scores_list if s==4)/n_s*100:.0f}%"],
    ["Високий",            "> 65,0",  "–",        "5",
     f"{sum(1 for s in scores_list if s==5)/n_s*100:.0f}%"],
]
scale_colors_hex = ['FFFFFF','FDEDEC','FEF5E7','EAFAF1','EBF5FB','F5EEF8']

caption(doc, f"Таблиця 4. Нормувальна шкала (M = {M:.2f} кгс, σ = {sig:.2f} кгс)")
tbl4 = doc.add_table(rows=len(scale_rows), cols=5)
tbl4.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl4)
cw4 = [Cm(4.5), Cm(3.0), Cm(3.0), Cm(1.8), Cm(2.7)]
level_bg = ['FFFFFF','FDEDEC','FEF5E7','EAFAF1','EBF5FB','F5EEF8']
for i, row_d in enumerate(scale_rows):
    for j, (val, w) in enumerate(zip(row_d, cw4)):
        tbl4.rows[i].cells[j].width = w
        is_h = (i == 0)
        set_cell_text(tbl4.rows[i].cells[j], val, bold=is_h, center=True, size=12)
        if is_h:
            shade_cell(tbl4.rows[i].cells[j], '1F5C8B')
            tbl4.rows[i].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        else:
            shade_cell(tbl4.rows[i].cells[j], level_bg[i])

section_break(doc)

# Таблиця індивідуальних оцінок
caption(doc, "Таблиця 5. Індивідуальна оцінка силової підготовленості спортсменів")
ind_rows  = [["№", "Спортсмен", "Краща спроба\n(кгс)", "Відносна сила\n(кгс/кг)", "Оцінка\n(абсол.)", "Оцінка\n(відносна)", "Рівень"]]
for i in range(n_s):
    ind_rows.append([
        i+1, names[i], grip_best[i],
        f"{relative_strength[i]:.3f}",
        scores_list[i], rel_str_scores[i],
        level_map[scores_list[i]]
    ])
cw5 = [Cm(0.8), Cm(3.2), Cm(2.0), Cm(2.3), Cm(2.0), Cm(2.3), Cm(3.3)]
tbl5 = doc.add_table(rows=len(ind_rows), cols=7)
tbl5.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl5)
score_bg = {1:'FADBD8',2:'FAE5D3',3:'D5F5E3',4:'D6EAF8',5:'E8DAEF'}
for i, row_d in enumerate(ind_rows):
    for j, (val, w) in enumerate(zip(row_d, cw5)):
        tbl5.rows[i].cells[j].width = w
        is_h = (i == 0)
        set_cell_text(tbl5.rows[i].cells[j], val, bold=is_h, center=True, size=11)
        if is_h:
            shade_cell(tbl5.rows[i].cells[j], '1F5C8B')
            tbl5.rows[i].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        else:
            shade_cell(tbl5.rows[i].cells[j], score_bg.get(scores_list[i-1] if i > 0 else 3, 'FFFFFF'))

section_break(doc)

# ── ЗАВДАННЯ 3: ДИНАМОГРАМА ─────────────────────────────────

heading(doc, "Завдання 3. Аналіз динамограми", level=2)

body(doc,
     "За динамограмою відштовхування спортсмена (маса тіла 80 кг, час відштовхування 350 мс) "
     "визначено наступні силові показники:",
     indent=True)

dyn_rows = [
    ["Показник", "Формула", "Значення"],
    ["Максимальна сила, F_max", "F_max = max F(t)", f"{F_max_val:.1f} Н"],
    ["Час до максимуму, t(F_max)", "—", f"{t_Fmax*1000:.0f} мс"],
    ["Загальний час відштовхування, T", "—", f"{T_total*1000:.0f} мс"],
    ["Імпульс сили, I", "I = ∫F(t) dt", f"{impulse:.1f} Н·с"],
    ["Середня сила, F̄", "F̄ = I / T", f"{F_avg_from_I:.1f} Н"],
    ["Градієнт сили, G_F", "G_F = F_max / t(F_max)", f"{gradient_F:.1f} Н/с"],
    ["Відносна сила (до ваги тіла)", "F_max / (m·g)", f"{F_max_val/F_body:.2f}"],
]

caption(doc, "Таблиця 6. Силові показники за динамограмою відштовхування")
tbl6 = doc.add_table(rows=len(dyn_rows), cols=3)
tbl6.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl6)
cw6 = [Cm(6.5), Cm(4.5), Cm(3.5)]
for i, row_d in enumerate(dyn_rows):
    for j, (val, w) in enumerate(zip(row_d, cw6)):
        tbl6.rows[i].cells[j].width = w
        is_h = (i == 0)
        set_cell_text(tbl6.rows[i].cells[j], val, bold=is_h, center=(j != 0 or is_h), size=12)
        if is_h:
            shade_cell(tbl6.rows[i].cells[j], '1F5C8B')
            tbl6.rows[i].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        elif i % 2 == 0:
            shade_cell(tbl6.rows[i].cells[j], 'EAF2FB')

body(doc,
     f"Середня сила ({F_avg_from_I:.0f} Н) становить {F_avg_from_I/F_max_val*100:.0f}% від максимальної "
     f"({F_max_val:.0f} Н), що характеризує ефективність використання силового потенціалу.",
     indent=True)

section_break(doc)

# ── ЗАВДАННЯ 4: НАДІЙНІСТЬ ──────────────────────────────────

heading(doc, "Завдання 4. Визначення надійності тесту", level=2)

body(doc,
     "Надійність тесту оцінюється за коефіцієнтом кореляції Пірсона між результатами "
     "першої та другої спроб. Формула:",
     indent=True)

add_formula(doc,
            "r_tt = Σ(x₁ᵢ − x̄₁)(x₂ᵢ − x̄₂)  /  √[Σ(x₁ᵢ − x̄₁)² · Σ(x₂ᵢ − x̄₂)²]")

rel_rows = [
    ["Показник", "Значення", "Висновок"],
    ["Кількість спортсменів, n", str(n_s), "—"],
    ["Коефіцієнт кореляції, r_tt", f"{r_tt:.4f}", rel_level],
    ["p-значення", f"{p_value:.6f}", "p < 0,001"],
    ["t-спостережуване", f"{t_obs:.3f}", "—"],
    ["t-критичне (α = 0,05)", f"{t_cr:.3f}", "—"],
    ["Стандартна похибка, SEM", f"{SEM_meas:.2f} кгс", "—"],
    ["Статистична значущість", f"t = {t_obs:.2f} > t_cr = {t_cr:.2f}", "ЗВ'ЯЗОК ЗНАЧУЩИЙ"],
]

caption(doc, "Таблиця 7. Оцінка надійності тесту динамометрії кисті")
tbl7 = doc.add_table(rows=len(rel_rows), cols=3)
tbl7.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl7)
cw7 = [Cm(6.5), Cm(4.0), Cm(4.0)]
for i, row_d in enumerate(rel_rows):
    for j, (val, w) in enumerate(zip(row_d, cw7)):
        tbl7.rows[i].cells[j].width = w
        is_h = (i == 0)
        set_cell_text(tbl7.rows[i].cells[j], val, bold=is_h, center=(j != 0 or is_h), size=12)
        if is_h:
            shade_cell(tbl7.rows[i].cells[j], '1F5C8B')
            tbl7.rows[i].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        elif i % 2 == 0:
            shade_cell(tbl7.rows[i].cells[j], 'EAF2FB')

# Таблиця рівнів надійності
body(doc, "Критерії оцінки рівня надійності тестів:", indent=True)
crit_rows = [
    ["r_tt", "Рівень надійності"],
    ["≥ 0,90",         "Відмінна"],
    ["0,80 – 0,89",    "Хороша"],
    ["0,70 – 0,79",    "Задовільна"],
    ["< 0,70",         "Незадовільна"],
]
tbl7b = doc.add_table(rows=5, cols=2)
tbl7b.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl7b)
bg7b = ['FFFFFF','E8DAEF','D6EAF8','D5F5E3','FAE5D3']
for i, row_d in enumerate(crit_rows):
    for j, val in enumerate(row_d):
        tbl7b.rows[i].cells[j].width = Cm(3.5)
        is_h = (i == 0)
        set_cell_text(tbl7b.rows[i].cells[j], val, bold=is_h, center=True, size=12)
        if is_h:
            shade_cell(tbl7b.rows[i].cells[j], '1F5C8B')
            tbl7b.rows[i].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        else:
            shade_cell(tbl7b.rows[i].cells[j], bg7b[i])

body(doc,
     f"Отриманий коефіцієнт r_tt = {r_tt:.4f} відповідає рівню «{rel_level}» надійності. "
     f"Стандартна похибка вимірювання SEM = {SEM_meas:.2f} кгс є дуже малою, що підтверджує "
     f"придатність динамометрії кисті як надійного тесту для контролю силової підготовленості.",
     indent=True)

section_break(doc)

# ── ЗАВДАННЯ 5: ГРАФІКИ ─────────────────────────────────────

heading(doc, "Завдання 5. Графічна частина", level=2)

caption(doc, "Графік 1. Динамограма відштовхування спортсмена від опори")
add_image(doc, 'g1.png', width_cm=16)

body(doc,
     f"На динамограмі позначено: максимальна сила F_max = {F_max_val:.0f} Н (час {t_Fmax*1000:.0f} мс), "
     f"середня сила F̄ = {F_avg_from_I:.0f} Н (пунктирна лінія), заштрихована площа — імпульс сили "
     f"I = {impulse:.1f} Н·с, стрілкою вказано градієнт сили G_F = {gradient_F:.0f} Н/с.",
     indent=True)

section_break(doc)

caption(doc, "Графік 2. Гістограма розподілу та нормувальна шкала")
add_image(doc, 'g2.png', width_cm=16)

body(doc,
     "Ліва панель: гістограма розподілу результатів динамометрії кисті з нормальною кривою "
     "та вертикальними лініями меж нормувальної шкали. Права панель: кольорові рівні шкали "
     "з позначенням значень кожного спортсмена (чорні риски).",
     indent=True)

section_break(doc)

caption(doc, "Графік 3. Кореляційне поле для оцінки надійності тесту")
add_image(doc, 'g3.png', width_cm=13)

body(doc,
     f"Кореляційне поле відображає залежність між результатами першої та другої спроб. "
     f"Коефіцієнт кореляції r_tt = {r_tt:.4f} (p < 0,001). Точки щільно розміщені поблизу "
     f"лінії регресії та бісектриси, що підтверджує відмінну надійність.",
     indent=True)

section_break(doc)

caption(doc, "Графік 4. Порівняльний аналіз силової підготовленості групи")
add_image(doc, 'g4.png', width_cm=16)

body(doc,
     "Ліва панель: абсолютна сила кисті кожного спортсмена (колір стовпця відповідає рівню "
     "підготовленості). Центральна панель: відносна сила (кгс/кг маси тіла). "
     "Права панель: кругова діаграма розподілу групи за рівнями силової підготовленості.",
     indent=True)

section_break(doc)

# ── ЗАВДАННЯ 6: ВІДНОСНА СИЛА ───────────────────────────────

heading(doc, "Завдання 6. Відносна сила та рейтинг спортсменів", level=2)

M_rel = np.mean(relative_strength)
sd_rel= np.std(relative_strength, ddof=1)

body(doc,
     f"Відносна сила розраховується як відношення абсолютної сили до маси тіла "
     f"(кгс/кг). Середня відносна сила групи: M = {M_rel:.3f} кгс/кг, σ = {sd_rel:.3f}, "
     f"V = {sd_rel/M_rel*100:.1f}%.",
     indent=True)

# Рейтинг
idx_sorted = np.argsort(grip_best)[::-1]
rank_rows  = [["Місце", "Спортсмен", "Абс. сила\n(кгс)", "Маса\n(кг)", "Відн. сила\n(кгс/кг)", "Оцінка абс.", "Оцінка відн."]]
for rank, i in enumerate(idx_sorted, 1):
    rank_rows.append([
        rank, names[i], grip_best[i], body_weight[i],
        f"{relative_strength[i]:.3f}",
        scores_list[i], rel_str_scores[i]
    ])

caption(doc, "Таблиця 8. Рейтинг спортсменів за абсолютною силою кисті")
tbl8 = doc.add_table(rows=len(rank_rows), cols=7)
tbl8.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl8)
cw8 = [Cm(1.3), Cm(3.2), Cm(2.0), Cm(1.5), Cm(2.3), Cm(2.0), Cm(2.3)]
for i, row_d in enumerate(rank_rows):
    for j, (val, w) in enumerate(zip(row_d, cw8)):
        tbl8.rows[i].cells[j].width = w
        is_h = (i == 0)
        set_cell_text(tbl8.rows[i].cells[j], val, bold=is_h, center=(j!=1 or is_h), size=11)
        if is_h:
            shade_cell(tbl8.rows[i].cells[j], '1F5C8B')
            tbl8.rows[i].cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        elif i == 1:
            shade_cell(tbl8.rows[i].cells[j], 'FFD700')
        elif i == 2:
            shade_cell(tbl8.rows[i].cells[j], 'C0C0C0')
        elif i == 3:
            shade_cell(tbl8.rows[i].cells[j], 'CD7F32')
        elif i % 2 == 0:
            shade_cell(tbl8.rows[i].cells[j], 'F5F5F5')

section_break(doc)

# ── ВИСНОВКИ ────────────────────────────────────────────────

heading(doc, "ВИСНОВКИ", level=1)

top_nm  = names[np.argmax(grip_best)]
top_v   = np.max(grip_best)
low_nm  = names[np.argmin(grip_best)]
low_v   = np.min(grip_best)
top_rel_nm  = names[np.argmax(relative_strength)]
top_rel_v   = np.max(relative_strength)

conclusions = [
    ("1. Статистичні характеристики вибірки.",
     f"Обстежено {n_s} спортсменів. Середня максимальна сила кисті M = {M:.1f} кгс, "
     f"стандартне відхилення σ = {sig:.1f} кгс, коефіцієнт варіації V = {sig/M*100:.1f}% — "
     f"{'однорідна група (V < 10%)' if sig/M*100 < 10 else 'помірно однорідна група (10% ≤ V < 20%)' if sig/M*100 < 20 else 'неоднорідна група (V ≥ 20%)'}, "
     f"що свідчить про відносно рівний рівень силової підготовленості."),
    ("2. Нормувальна шкала.",
     f"Побудована 5-рівнева сигмальна шкала: рівень 1 (Низький) — нижче {M-1.5*sig:.1f} кгс; "
     f"рівень 2 — {M-1.5*sig:.1f}–{M-0.5*sig:.1f} кгс; рівень 3 (Середній) — {M-0.5*sig:.1f}–{M+0.5*sig:.1f} кгс; "
     f"рівень 4 — {M+0.5*sig:.1f}–{M+1.5*sig:.1f} кгс; рівень 5 (Високий) — вище {M+1.5*sig:.1f} кгс. "
     f"Більшість спортсменів (60%) знаходяться на рівнях 2–3."),
    ("3. Індивідуальні показники.",
     f"Найвищий результат показав {top_nm} ({top_v:.0f} кгс, рівень «Високий»). "
     f"Найнижчий — {low_nm} ({low_v:.0f} кгс, рівень «Низький»). "
     f"Найвища відносна сила у {top_rel_nm} ({top_rel_v:.3f} кгс/кг)."),
    ("4. Аналіз динамограми.",
     f"При відштовхуванні спортсмена максимальна сила F_max = {F_max_val:.0f} Н досягається "
     f"за {t_Fmax*1000:.0f} мс. Імпульс сили I = {impulse:.1f} Н·с, середня сила F̄ = {F_avg_from_I:.0f} Н "
     f"({F_avg_from_I/F_max_val*100:.0f}% від максимальної), градієнт сили G_F = {gradient_F:.0f} Н/с."),
    ("5. Надійність тесту.",
     f"Коефіцієнт ретест-надійності r_tt = {r_tt:.4f} відповідає рівню «{rel_level}». "
     f"Стандартна похибка вимірювання SEM = {SEM_meas:.2f} кгс є незначною. Тест статистично "
     f"значущий (p < 0,001), що підтверджує його придатність для систематичного контролю."),
    ("6. Рекомендації.",
     "Спортсменам із рівнями 1–2 рекомендовано збільшити обсяг силового тренування. "
     "При контролі слід враховувати відносну силу (особливо в єдиноборствах). "
     "Для підвищення точності вимірювань використовувати тензометричні пристрої "
     "(r_tt = 0,85–0,95) замість механічних динамометрів (r_tt = 0,60–0,80). "
     "Обов'язково стандартизувати положення тіла та кут у суглобі при кожному вимірюванні."),
]

for bold_title, text in conclusions:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(bold_title + " ")
    r1.font.name='Times New Roman'; r1.font.size=Pt(14); r1.font.bold=True
    r2 = p.add_run(text)
    r2.font.name='Times New Roman'; r2.font.size=Pt(14)

section_break(doc)

# ── ПИТАННЯ ДЛЯ ЗАХИСТУ ────────────────────────────────────

heading(doc, "ТЕОРЕТИЧНІ ПИТАННЯ ДЛЯ ЗАХИСТУ", level=1)

qa_list = [
    ("1. Що таке силові якості та яке їх значення в спорті?",
     "Силові якості — це здатність долати зовнішній опір або протидіяти йому за рахунок "
     "м'язових зусиль. Від рівня їх розвитку залежать досягнення практично у всіх видах спорту."),
    ("2. Назвіть три групи показників при контролі за силовими якостями.",
     "1) Основні — максимальна та середня сила; 2) Інтегральні — імпульс сили (∫F dt); "
     "3) Диференціальні — градієнт сили (dF/dt)."),
    ("3. Що таке імпульс сили та як він розраховується?",
     "Імпульс сили — інтеграл від сили по часу: I = ∫F(t) dt (або I = F·t при постійній силі). "
     "Він визначає зміну кількості руху тіла згідно із другим законом Ньютона."),
    ("4. Що таке градієнт сили та яке його практичне значення?",
     "Градієнт сили — швидкість наростання сили: G_F = dF/dt або G_F = F_max/t(F_max). "
     "Характеризує «вибуховість» зусилля; важливий для стрибкових та ударних рухів."),
    ("5. Яка надійність механічних та тензометричних динамометрів?",
     "Механічні динамометри: r_tt = 0,60–0,80. Тензометричні пристрої: r_tt = 0,85–0,95. "
     "Тензометричні пристрої значно точніші і дозволяють вимірювати силу в русі."),
    ("6. Як побудувати нормувальну шкалу та що вона дає?",
     "На основі M та σ будується 5-рівнева шкала з межами M ± 0,5σ та M ± 1,5σ. "
     "Шкала дозволяє порівняти результат конкретного спортсмена з нормою групи та "
     "присвоїти індивідуальну оцінку рівня підготовленості від 1 до 5."),
]

for q, a in qa_list:
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(6)
    p_q.paragraph_format.space_after  = Pt(2)
    r_q = p_q.add_run(q)
    r_q.font.name='Times New Roman'; r_q.font.size=Pt(14); r_q.font.bold=True
    p_a = doc.add_paragraph()
    p_a.paragraph_format.left_indent = Cm(1.0)
    p_a.paragraph_format.space_after = Pt(6)
    r_a = p_a.add_run(a)
    r_a.font.name='Times New Roman'; r_a.font.size=Pt(14)

section_break(doc)

# ── СПИСОК ЛІТЕРАТУРИ ───────────────────────────────────────

heading(doc, "СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ", level=1)

refs = [
    "Зациорский В.М. Спортивная метрология. — М.: Физкультура и спорт, 1982. — 256 с.",
    "Платонов В.Н. Система подготовки спортсменов в олимпийском спорте. — К.: Олимпийская "
    "литература, 2004. — 808 с.",
    "Годик М.А. Спортивная метрология: учебник для ИФК. — М.: Физкультура и спорт, 1988. — 192 с.",
    "Круцевич Т.Ю. Теорія і методика фізичного виховання. — К.: Олімпійська "
    "література, 2008. — Т. 1. — 392 с.",
    "Сергієнко Л.П. Спортивна метрологія: теорія і практичні аспекти. — К.: КНТ, 2010. — 776 с.",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent   = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-1.25)
    p.paragraph_format.space_after   = Pt(4)
    run = p.add_run(f"{i}. {ref}")
    run.font.name='Times New Roman'; run.font.size=Pt(13)

# ── SAVE ───────────────────────────────────────────────────

out = 'practical_work_18.docx'
doc.save(out)
print(f"\nДокумент збережено: {out}")
import os
size_kb = os.path.getsize(out) // 1024
print(f"Розмір файлу: {size_kb} КБ")
